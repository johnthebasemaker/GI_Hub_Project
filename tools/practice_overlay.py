#!/usr/bin/env python3
"""
tools/practice_overlay.py — what a Practice SANDBOX needs that a tutorial VIDEO
does not, written on top of the Phase 12 synthetic dataset (rule 17, slice S2).

    GI_INSTANCE=training GI_DOTENV=0 \\
    DATABASE_URL=postgresql+asyncpg://gi_training@127.0.0.1:5433/gihub_seed_training \\
    .venv/bin/python tools/practice_overlay.py

Normally run by `tools/practice_db.py build`, never by hand.

⚠️ WHY THIS IS NOT AN EDIT TO `make_tutorial_db.py`. That file's output is
pinned by ruling P12-5: its `DATASET_VERSION` is written into every rendered
tutorial's manifest, so adding trainee accounts or seeded queues THERE would
change the dataset sixty videos were recorded against. This overlay runs after
`cutover_migrate` has loaded the fixture, touches only the Postgres copy, and
leaves the tutorial dataset byte-identical. The side benefit is deliberate: a
trainee who watches "Return Stock" and then opens Practice finds *Tutorial Demo
Gasket Set / 899001* exactly where the video showed it.

⚠️ IT RUNS AS A PRACTICE PROCESS. `GI_INSTANCE=training` against the SEED
database (`gihub_seed_training`, which ends `_training` for exactly this
reason), so rule 17's boot check applies to it like any other Practice process:
pointed at Live, it refuses before it has an engine.

WHAT IT ADDS
  · one SHARED account per role in `auth.ROLE_META` (`practice.<role>`), with
    the scope each role's registration rule demands — ruling Q4. A role added to
    ROLE_META without an account here fails suite TR-10 (rule 13's pattern).
  · the five tutorial harness logins are REMOVED: `admin / admin2026` is a
    password published in the repo, and in a sandbox anybody can reach it would
    be an admin account nobody chose to hand out.
  · work waiting in every approval queue, staged through the REAL services, so
    an HOD trainee opens Approvals and finds something to approve.
  · 2FA mandate off (Q4), phones replaced by a non-routable number (V7's second
    wall), and the dataset stamps `GET /instance` reports.
"""
from __future__ import annotations

import asyncio
import datetime as _dt
import os
import pathlib
import secrets
import sys

_ROOT = pathlib.Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))
os.environ.setdefault("GI_DOTENV", "0")

# Bumped when the OVERLAY's content changes. Reported with the tutorial
# fixture's own DATASET_VERSION as "<fixture>.<overlay>".
OVERLAY_VERSION = 11  # 2 = Phase 15e two-note job · 3 = Phase 16 lots & expiry (2026-10-01)
                      # · 4 = Phase 18 return desk loans (2026-10-03)
                      # · 5 = Phase 18 reorder-signal trio (2026-10-03)
                      # · 6 = Phase 19 accepted minimum · per-site / global POs ·
                      #       a partly returned loan (2026-10-04)
                      # · 7 = Phase 20 a week of Surface Shield jobs in every
                      #       status — the daily log and the bulk flows (2026-10-04)
                      # · 8 = Phase 21a an empty lot the dashboard must skip and
                      #       workbook lot problems with a sheet and row (2026-10-06)
                      # · 9 = Phase 21d learned OCR names — the paste lane shows
                      #       green (learned), gold (suggested) and red (2026-10-06)
                      # · 10 = Phase 21f the self-driving demo's own tank
                      #       (DEMO-TANK-1, two READY days) (2026-10-06)
                      # · 11 = Phase 22 the Drive features on Practice data — a DN
                      #       to open, a WD delivery, two certificates (one for QC),
                      #       a request, Day/Night preparers, a learned tank, a Night
                      #       paper to Compare; the queue seed made idempotent (2026-10-08)

SITE = "CNCEC"
WAREHOUSE = "WH-01"
WBS = "WBS-9001"
# The tutorial harness accounts (make_tutorial_db.HARNESS_USERS). Removed here.
HARNESS_USERNAMES = ("admin", "hod", "supervisor", "worker", "Logistics")
# A number no carrier routes. Outbound is already impossible in Practice; this
# is so a copied-out row can never become a real text to a real person.
NON_ROUTABLE_PHONE = "+000000000000"
DEFAULT_PASSWORD = "Practice@2026"

# (username, role, Site_ID, Warehouse_ID). Scopes follow auth.py's
# _SCOPED_ / _UNSCOPED_ / _DUAL_SCOPE_REG_ROLES exactly — TR-10 asserts it.
ACCOUNTS: tuple[tuple[str, str, str, str | None], ...] = (
    ("practice.admin",       "admin",          "",   None),
    ("practice.logistics",   "logistics",      "",   None),
    ("practice.auditor",     "auditor",        "",   None),
    ("practice.hod",         "hod",            SITE, None),
    ("practice.qchod",       "qc_hod",         "",   None),
    ("practice.warehouse",   "warehouse_user", "",   WAREHOUSE),
    ("practice.supervisor",  "supervisor",     SITE, None),
    ("practice.qc",          "qc",             SITE, None),
    ("practice.storekeeper", "store_keeper",   SITE, None),
)


def practice_password() -> str:
    return os.environ.get("PRACTICE_PASSWORD", "").strip() or DEFAULT_PASSWORD


def admin_password() -> tuple[str, bool]:
    """The Practice ADMIN is the one account that can wipe everybody's work,
    so it does not get the published password unless the operator says so.
    Returns (password, generated?)."""
    v = os.environ.get("PRACTICE_ADMIN_PASSWORD", "").strip()
    if v:
        return v, False
    return "Pa-" + secrets.token_urlsafe(12), True


async def seed_accounts(session) -> dict:
    """Idempotent: re-running converges on the same nine rows. Does NOT commit."""
    import bcrypt
    from sqlalchemy import delete, select, update
    from sqlalchemy.dialects.postgresql import insert as pg_insert

    from backend import models
    users = models.Base.metadata.tables["users"]
    await session.execute(delete(users).where(users.c["username"].in_(HARNESS_USERNAMES)))
    pw = practice_password()
    admin_pw, generated = admin_password()
    for uname, role, site, wh in ACCOUNTS:
        raw = admin_pw if role == "admin" else pw
        h = bcrypt.hashpw(raw.encode(), bcrypt.gensalt(rounds=10)).decode()
        vals = {"password_hash": h, "role": role, "Site_ID": site,
                "Warehouse_ID": wh, "Phone_Number": None, "email": None,
                "totp_secret": None, "totp_enabled": 0}
        await session.execute(pg_insert(users).values(username=uname, **vals)
                              .on_conflict_do_update(index_elements=["username"],
                                                     set_=vals))
    n = (await session.execute(select(users.c["username"]).where(
        users.c["username"].like("practice.%")))).all()
    return {"accounts": len(n), "admin_password": admin_pw if generated else None}


async def seed_settings(session, *, fixture_version: int) -> None:
    """Idempotent. Does NOT commit."""
    from sqlalchemy import update
    from sqlalchemy.dialects.postgresql import insert as pg_insert

    from backend import models
    t = models.Base.metadata.tables["app_settings"]
    emp = models.Base.metadata.tables["employees"]
    now = _dt.datetime.now(_dt.timezone.utc).isoformat(timespec="seconds")
    for k, v in (("mfa_required_roles", ""),              # ruling Q4
                 ("maintenance_mode", "0"),
                 ("practice_dataset_version", f"{fixture_version}.{OVERLAY_VERSION}"),
                 ("practice_seeded_at", now)):
        await session.execute(pg_insert(t).values(key=k, value=v)
                              .on_conflict_do_update(index_elements=["key"],
                                                     set_={"value": v}))
    await session.execute(update(emp).values(Phone_Number=NON_ROUTABLE_PHONE))


async def _set(session, key: str, value: str | None) -> None:
    from sqlalchemy import delete
    from sqlalchemy.dialects.postgresql import insert as pg_insert

    from backend import models
    t = models.Base.metadata.tables["app_settings"]
    if value is None:
        await session.execute(delete(t).where(t.c["key"] == key))
    else:
        await session.execute(pg_insert(t).values(key=key, value=value)
                              .on_conflict_do_update(index_elements=["key"],
                                                     set_={"value": value}))


async def seed_queues() -> dict:
    """Stage real work through the REAL app, so every side effect a trainee
    will see — the HOD's bell, the notification text, the pending row's shape —
    is the one Live produces. Returns counts per queue."""
    from httpx import ASGITransport, AsyncClient
    from sqlalchemy import func, select

    from backend import models
    from backend.api.db import SessionLocal
    from backend.api.main import app
    from backend.api.services import ledger, procurement

    md = models.Base.metadata
    today = _dt.date.today().isoformat()
    pw = practice_password()
    out: dict[str, int] = {}

    # The supporting-document gate is ON in Live and stays ON in Practice
    # (parity). It is lifted only for the seconds it takes to stage the seed —
    # the trainee's own entries meet it exactly as in Live.
    async with SessionLocal() as s:
        t = md.tables["app_settings"]
        prev = (await s.execute(select(t.c["value"]).where(
            t.c["key"] == "require_entry_documents"))).scalar_one_or_none()
        await _set(s, "require_entry_documents", "0")
        await s.commit()
    try:
        async with AsyncClient(transport=ASGITransport(app=app),
                               base_url="http://practice-overlay") as ac:
            r = await ac.post("/auth/login", json={"username": "practice.storekeeper",
                                                   "password": pw})
            r.raise_for_status()
            H = {"Authorization": f"Bearer {r.json()['access_token']}",
                 "X-GI-Instance": "training"}
            async with SessionLocal() as s:
                inv = md.tables["inventory"]
                saps = [x[0] for x in (await s.execute(
                    select(inv.c["SAP_Code"]).where(inv.c["Site_ID"] == SITE)
                    .order_by(inv.c["SAP_Code"]).limit(6))).all()]
            receipts = [("899001", 24.0), (saps[0], 40.0), (saps[1], 12.0)]
            # Phase 22f — IDEMPOTENT: a second run used to add another 3 + 3 + 1
            # rows and two PRs (12 seed receipts after four runs). Seeded once;
            # earlier duplicates are trimmed back to one run's worth.
            async with SessionLocal() as s:
                already = await trim_seed_duplicates(s)
                await s.commit()
            if already:
                receipts = []
            for sap, qty in receipts:
                rr = await ac.post("/entry/receipts", headers=H, json={
                    "Date": today, "SAP_Code": sap, "Quantity": qty, "Site_ID": SITE,
                    "Supplier": "Halcyon Industrial Supply", "wbs": WBS,
                    "Remarks": "Practice seed — approve or reject me"})
                if rr.status_code != 201:
                    print(f"  ⚠️  receipt {sap}: {rr.status_code} {rr.text[:160]}")
            issues = [] if already else [("899001", 6.0, "Aria Bellweather"),
                                         (saps[2], 3.0, "Tomas Halversen"),
                                         (saps[3], 2.0, "Nadia Okonjo")]
            for sap, qty, who in issues:
                ri = await ac.post("/entry/consumption", headers=H, json={
                    "Date": today, "SAP_Code": sap, "Quantity": qty, "Site_ID": SITE,
                    "Work_Type": "Maintenance", "Tank_No": "Tank 1", "Issued_To": who,
                    "wbs": WBS, "Remarks": "Practice seed — approve or reject me"})
                if ri.status_code != 201:
                    print(f"  ⚠️  issue {sap}: {ri.status_code} {ri.text[:160]}")
            try:
                out["garnet_jobs"] = await _seed_garnet(today)
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  garnet example skipped: {type(e).__name__}: {str(e)[:160]}")
            try:
                out["note_jobs"] = await seed_job_notes(today)
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  job-note example skipped: {type(e).__name__}: {str(e)[:160]}")
            try:
                out["log_jobs"] = await seed_daily_log_jobs(today)
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  daily-log example skipped: {type(e).__name__}: {str(e)[:160]}")
            try:
                out["demo_jobs"] = await seed_demo_jobs(today)
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  demo-tank example skipped: {type(e).__name__}: {str(e)[:160]}")
            try:
                out["lots"] = await seed_lots(today)
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  lot example skipped: {type(e).__name__}: {str(e)[:160]}")
            try:
                out["lot_problems"] = await seed_lot_problems(today)
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  lot-problem example skipped: {type(e).__name__}: {str(e)[:160]}")
            try:
                out["ocr_names"] = await seed_ocr_names()
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  learned-name example skipped: {type(e).__name__}: {str(e)[:160]}")
            try:
                out["reorder"] = await seed_reorder_examples()
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  reorder example skipped: {type(e).__name__}: {str(e)[:160]}")
            try:
                out["accepted_min"] = await seed_accepted_minimum()
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  accepted-minimum example skipped: {type(e).__name__}: {str(e)[:160]}")
            try:
                out["on_order"] = await seed_on_order_examples()
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  on-order example skipped: {type(e).__name__}: {str(e)[:160]}")
            try:
                out["drive"] = await seed_drive_examples(today)
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  Drive examples skipped: {type(e).__name__}: {str(e)[:160]}")
            try:
                out["loans"] = await seed_returnables()
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  loan example skipped: {type(e).__name__}: {str(e)[:160]}")
            try:
                out["partial_loan"] = await seed_partial_return()
            except Exception as e:  # noqa: BLE001 — an example, never a blocker
                print(f"  ⚠️  partial-return example skipped: {type(e).__name__}: {str(e)[:160]}")

        async with SessionLocal() as s:
            draft: dict = {}
            sent: dict = {}
            if not already:
                await ledger.stage_return(s, username="practice.storekeeper", data={
                    "Site_ID": SITE, "SAP_Code": "899001", "Quantity": 2.0,
                    "Reason": "Surplus to job", "Lot_Number": "LOT-TUT-001",
                    "Remarks": "Practice seed"})
                await s.commit()
                draft = await procurement.create_pr(
                    s, username="practice.hod", site_id=SITE,
                    lines=[{"SAP_Code": saps[4], "Requested_Qty": 50},
                           {"SAP_Code": saps[5], "Requested_Qty": 20}],
                    notes="Practice seed — a draft for the HOD to submit")
                sent = await procurement.create_pr(
                    s, username="practice.hod", site_id=SITE,
                    lines=[{"SAP_Code": "899001", "Requested_Qty": 100}],
                    notes="Practice seed — already with Logistics")
                await s.commit()
                if sent.get("pr_number"):
                    await procurement.submit_pr(s, username="practice.hod",
                                                pr_number=sent["pr_number"], site_id=SITE)
                    await s.commit()
            for name in ("pending_receipts", "pending_issues", "pending_returns"):
                tt = md.tables[name]
                out[name] = (await s.execute(select(func.count()).select_from(tt).where(
                    tt.c["status"] == ledger.PENDING))).scalar_one()
            pm = md.tables["pr_master"]
            out["pr_draft"] = (await s.execute(select(func.count(func.distinct(
                pm.c["PR_Number"]))).where(pm.c["logistics_status"] == "site_draft"))).scalar_one()
            if draft.get("error") or sent.get("error"):
                print(f"  ⚠️  PR seed: {draft.get('error') or sent.get('error')}")
    finally:
        async with SessionLocal() as s:
            await _set(s, "require_entry_documents", prev)
            await s.commit()
    return out


_SEED_REMARK = "Practice seed — approve or reject me"


async def trim_seed_duplicates(s) -> bool:
    """Phase 22f: True when the queue seed already ran (so it is not staged
    again). Rows a non-idempotent older overlay added on every run are trimmed
    back to ONE run's worth — the oldest kept: 3 receipts, 3 issues, 1 return."""
    from sqlalchemy import text
    seen = (await s.execute(text(
        'SELECT COUNT(*) FROM pending_receipts WHERE "Remarks" = :r'), {"r": _SEED_REMARK})).scalar() or 0
    if not seen:
        return False
    for table, remark, keep in (("pending_receipts", _SEED_REMARK, 3),
                                ("pending_issues", _SEED_REMARK, 3),
                                ("pending_returns", "Practice seed", 1)):
        res = await s.execute(text(f"""
            DELETE FROM {table} WHERE "Remarks" = :r AND id NOT IN (
                SELECT id FROM {table} WHERE "Remarks" = :r ORDER BY id LIMIT :k)"""),
            {"r": remark, "k": keep})
        if res.rowcount:
            print(f"  trimmed {res.rowcount} duplicate seed row(s) from {table}")
    return True


GARNET_SAP = "899970"   # synthetic, Practice-only (P12-5)


async def _seed_garnet(today: str) -> int:
    """Phase 15d — one Garnet draw on a steel tank, so the supervisor's queue
    has a Surface-prep card to answer Old / New on, and the HOD a benchmark to
    read.

    Practice-only data, so it lives here and never in make_tutorial_db (P12-5).
    Never fatal: a Practice without this example is still a Practice.
    """
    from sqlalchemy import text

    from backend.api.db import SessionLocal
    from backend.api.services import prep as PR
    from backend.api.services import quality

    async with SessionLocal() as s:
        cat = await quality.controlled_category(s)
        tag = None
        for t, ty, su in (await s.execute(text(
                'SELECT "Equipment_Tag_No", "Type", "Substrate" FROM sme_equipment '
                'WHERE "Site_ID" = :s ORDER BY "Equipment_Tag_No"'), {"s": SITE})).all():
            if PR.code_for(ty, su) == "ESC2":
                tag = t
                break
        if tag is None:
            print("  ⚠️  garnet: no steel / vessel equipment to put the example on")
            return 0
        await s.execute(text(
            'INSERT INTO inventory ("SAP_Code", "Material_Code", "Equipment_Description", '
            '"Category", "UOM", "Site_ID", "Unit_Size", "Base_UOM", "Opening_Stock") VALUES '
            "(:p, 'MAT-899970', 'PRACTICE GARNET 30/60 MESH', :c, 'TON', :site, 1000, 'KG', 10) "
            'ON CONFLICT ("SAP_Code") DO NOTHING'), {"p": GARNET_SAP, "c": cat, "site": SITE})
        for code, name, rate in (("ESC1", "Blasting Civil Floor & Wall", 18),
                                 ("ESC2", "Blasting Steel Surface", 20)):
            await s.execute(text(
                'INSERT INTO sme_recipe ("Lining_System_Code", "Execution_Sub_Activity_Code", '
                '"Lining_System_Name", "Material_Code", "SAP_Code", "Material_Name", '
                '"Material_Description", "UOM", "For_1_SQM") VALUES (:c, :c, :n, '
                "'MAT-899970', :p, 'PRACTICE GARNET 30/60 MESH', 'Garnet', 'KG', :r)"),
                {"c": code, "n": name, "p": GARNET_SAP, "r": rate})
        await s.commit()
    # ⚠️ A LEDGER ROW, not a staged issue: issuing a Surface Shield needs a QC
    # release and an MTC on file, a chain this example has no business
    # faking. Live's Garnet draws arrive the same way — from the Excel sync —
    # and the queue sweeps every ledger row whatever its route (ruling Q13-6).
    async with SessionLocal() as s:
        await s.execute(text(
            'INSERT INTO consumption ("Date", "SAP_Code", "Quantity", "Site_ID", "Tank_No", '
            '"Work_Type", "Issued_To", "Remarks") VALUES (:d, :p, 2, :site, :t, '
            "'Blasting', 'Tomas Halversen', 'Practice seed — Garnet — Shell - 90 SQM Done')"),
            {"d": today, "p": GARNET_SAP, "site": SITE, "t": tag})
        await s.commit()
    return 1


JOB_TAG = "PRACTICE-TK-01"
JOB_CODE = "PRL1"
JOB_SAPS = (("899971", "PRACTICE PU PRIMER", 0.35), ("899972", "PRACTICE PU TOPCOAT", 0.6))


async def seed_job_notes(today: str) -> int:
    """Phase 15e — a lining job whose store-keeper remarks fill the card.

    One Practice tank with its own two-material system, and ONE day of draws
    carrying TWO different notes ("Floor - 12.5 SQM Done", "Sump Wall - 4.2
    SQM Done"), so a supervisor trainee sees the note buttons, the pre-filled
    area / part / remark, and submits the two as two jobs; the HOD then sees the
    remark on the approval card. The fixture's own recipes name no stock item
    (their SAP_Code is blank), so no lining job could be practised without this.

    Idempotent (keyed on the synthetic SAPs / tag) and never fatal. Practice-only
    data (P12-5). Rule 17's standing order: a Live feature gets its example here.
    """
    from sqlalchemy import text

    from backend.api.db import SessionLocal
    from backend.api.services import quality

    async with SessionLocal() as s:
        if (await s.execute(text(
                'SELECT 1 FROM consumption WHERE "Tank_No" = :t LIMIT 1'), {"t": JOB_TAG})).first():
            return 0
        cat = await quality.controlled_category(s)
        for sap, name, rate in JOB_SAPS:
            await s.execute(text(
                'INSERT INTO inventory ("SAP_Code", "Material_Code", "Equipment_Description", '
                '"Category", "UOM", "Site_ID", "Opening_Stock") VALUES '
                "(:p, :m, :n, :c, 'KG', :site, 200) ON CONFLICT (\"SAP_Code\") DO NOTHING"),
                {"p": sap, "m": f"MAT-{sap}", "n": name, "c": cat, "site": SITE})
            await s.execute(text(
                'INSERT INTO sme_recipe ("Lining_System_Code", "Execution_Sub_Activity_Code", '
                '"Lining_System_Name", "Material_Code", "SAP_Code", "Material_Name", '
                '"Material_Description", "UOM", "For_1_SQM") VALUES (:c, :c, '
                "'Practice PU lining', :m, :p, :n, :n, 'KG', :r)"),
                {"c": JOB_CODE, "m": f"MAT-{sap}", "p": sap, "n": name, "r": rate})
        await s.execute(text(
            'INSERT INTO sme_equipment ("Site_ID", "Equipment_Tag_No", "Name", "Type", '
            '"Substrate", "Lining_System_Code", "Surface_Area_SQM", "Equipment_Total_SQM") '
            "VALUES (:site, :t, 'Practice sump tank', 'CV', 'CONCRETE SUBSTRATE', :c, 120, 120)"),
            {"site": SITE, "t": JOB_TAG, "c": JOB_CODE})
        await s.execute(text(
            'INSERT INTO sme_sqm_progress ("Site_ID", "Equipment_Tag_No", "Lining_System_Code", '
            '"Original_SQM", "Done_SQM") VALUES (:site, :t, :c, 120, 0)'),
            {"site": SITE, "t": JOB_TAG, "c": JOB_CODE})
        # ⚠️ LEDGER ROWS, like the Garnet example: an issue of a Surface Shield
        # needs a QC release and an MTC this example has no business faking.
        for sap, qty, remark in (("899971", 4.4, "Floor - 12.5 SQM Done"),
                                 ("899972", 7.5, "Floor - 12.5 SQM Done"),
                                 ("899971", 1.5, "Sump Wall - 4.2 SQM Done"),
                                 ("899972", 2.5, "Sump Wall - 4.2 SQM Done")):
            await s.execute(text(
                'INSERT INTO consumption ("Date", "SAP_Code", "Quantity", "Site_ID", '
                '"Tank_No", "Work_Type", "Issued_To", "Remarks") VALUES '
                "(:d, :p, :q, :site, :t, 'Lining', 'Aria Bellweather', :r)"),
                {"d": today, "p": sap, "q": qty, "site": SITE, "t": JOB_TAG, "r": remark})
        await s.commit()
    return 2


async def seed_daily_log_jobs(today: str) -> int:
    """Phase 20 — a week of Surface Shield jobs, one of each status (rule 17g).

    On the Practice sump tank (PRACTICE-TK-01, PRL1: primer 0.35 + topcoat 0.6
    KG/m²), filed and decided THROUGH THE REAL SERVICES (`sme_groups.submit` /
    `decide_group`), so every figure is what Live would hold:

      D-6  "Floor - 10 SQM Done"          approved
      D-5  "Bottom Wall - 8 SQM Done"     approved, the HOD corrected it to 7.5
                                          with a justification (the remark and
                                          the approved SQM differ — the log says so)
      D-4  "Dyke wall patch work"         rejected: no area in the remark
      D-3  "Floor - 6 SQM Done"           pending
      D-2  "Sump Wall - 5 SQM Done"       pending, HIGH variance (primer +71 %)
      D-1  "Coving - 4 SQM Done"          not yet filed, and READY (the remark
                                          gives the SQM): the bulk-submit example

    Idempotent and never fatal; Practice-only data (P12-5)."""
    import datetime as _d

    from sqlalchemy import text

    from backend.api.db import SessionLocal
    from backend.api.services import sme_groups as G

    t0 = _d.date.fromisoformat(today)
    day = lambda n: (t0 - _d.timedelta(days=n)).isoformat()      # noqa: E731
    jobs = (  # days back, remark, primer KG, topcoat KG, filed SQM, decision
        (6, "Floor - 10 SQM Done", 3.5, 6.0, 10.0, "approve"),
        (5, "Bottom Wall - 8 SQM Done", 2.8, 4.8, 8.0, "approve-edit"),
        (4, "Dyke wall patch work", 1.0, 1.8, 3.0, "reject"),
        (3, "Floor - 6 SQM Done", 2.1, 3.6, 6.0, "pending"),
        (2, "Sump Wall - 5 SQM Done", 3.0, 3.0, 5.0, "pending"),
        (1, "Coving - 4 SQM Done", 1.4, 2.4, None, "unfiled"),
    )
    async with SessionLocal() as s:
        if (await s.execute(text(
                'SELECT 1 FROM consumption WHERE "Tank_No" = :t AND "Remarks" = :r LIMIT 1'),
                {"t": JOB_TAG, "r": jobs[0][1]})).first():
            return 0
        if not (await s.execute(text(
                'SELECT 1 FROM sme_equipment WHERE "Equipment_Tag_No" = :t'), {"t": JOB_TAG})).first():
            return 0        # seed_job_notes builds the tank; nothing to file against
        made = 0
        for back, remark, primer, top, sqm, what in jobs:
            ids = []
            for sap, qty in ((JOB_SAPS[0][0], primer), (JOB_SAPS[1][0], top)):
                ids.append((await s.execute(text(
                    'INSERT INTO consumption ("Date", "SAP_Code", "Quantity", "Site_ID", '
                    '"Tank_No", "Work_Type", "Issued_To", "Remarks") VALUES '
                    "(:d, :p, :q, :site, :t, 'Lining', 'Aria Bellweather', :r) RETURNING id"),
                    {"d": day(back), "p": sap, "q": qty, "site": SITE, "t": JOB_TAG,
                     "r": remark})).scalar_one())
            await s.commit()
            if sqm is None:
                made += 1
                continue
            async with s.begin():
                out = await G.submit(s, site_id=SITE, work_date=day(back), tag=JOB_TAG,
                                     code=JOB_CODE, sqm=sqm, consumption_ids=ids,
                                     notes=remark, username="practice.supervisor")
            gid = out["group_id"]
            if what.startswith("approve"):
                edit = what == "approve-edit"
                async with s.begin():
                    await G.decide_group(
                        s, group_id=gid, approve=True,
                        edits={"SQM_Completed": 7.5} if edit else None,
                        justification=("Re-measured on site: 7.5 m2 - the remark "
                                       "rounded up." if edit else ""),
                        reject_reason="", username="practice.hod", site_id=SITE)
            elif what == "reject":
                async with s.begin():
                    await G.decide_group(
                        s, group_id=gid, approve=False, edits=None, justification="",
                        reject_reason="No area in the remark - measure it and file again.",
                        username="practice.hod", site_id=SITE)
            made += 1
    return made


DEMO_TAG = "DEMO-TANK-1"


async def seed_demo_jobs(today: str) -> int:
    """Phase 21f — the self-driving demo's OWN tank (flow 2, `ss-bulk`).

    DEMO-TANK-1 on the Practice lining system (PRL1), with two days of
    store-keeper draws whose remarks give the area — so both cards are READY
    and the demo can select-all, submit, and have the HOD approve them. The
    draws are NOT tagged (they are the store keeper's, and they stay); the
    JOBS the demo files are on a DEMO- tank, so Reset demo data
    (`practice_demo.reset_demo_data`) un-files them and takes their m² back —
    the tank is ready for the next demo. Its own tank, so a trainee's work on
    PRACTICE-TK-01 is never filed by a demo. Idempotent; Practice-only (P12-5).
    """
    import datetime as _d

    from sqlalchemy import text

    from backend.api.db import SessionLocal

    t0 = _d.date.fromisoformat(today)
    day = lambda n: (t0 - _d.timedelta(days=n)).isoformat()      # noqa: E731
    async with SessionLocal() as s:
        if (await s.execute(text(
                'SELECT 1 FROM sme_equipment WHERE "Equipment_Tag_No" = :t'), {"t": DEMO_TAG})).first():
            return 0
        if not (await s.execute(text(
                'SELECT 1 FROM sme_recipe WHERE "Lining_System_Code" = :c LIMIT 1'), {"c": JOB_CODE})).first():
            return 0        # seed_job_notes builds the system; nothing to file against
        await s.execute(text(
            'INSERT INTO sme_equipment ("Site_ID", "Equipment_Tag_No", "Name", "Type", '
            '"Substrate", "Lining_System_Code", "Surface_Area_SQM", "Equipment_Total_SQM") '
            "VALUES (:site, :t, 'Demo tank (self-driving demo)', 'CV', 'CONCRETE SUBSTRATE', "
            ":c, 80, 80)"), {"site": SITE, "t": DEMO_TAG, "c": JOB_CODE})
        await s.execute(text(
            'INSERT INTO sme_sqm_progress ("Site_ID", "Equipment_Tag_No", "Lining_System_Code", '
            '"Original_SQM", "Done_SQM") VALUES (:site, :t, :c, 80, 0)'),
            {"site": SITE, "t": DEMO_TAG, "c": JOB_CODE})
        # recipe PRL1: primer 0.35 + topcoat 0.6 KG/m² — both days on the recipe
        for back, remark, primer, top in ((3, "Wall - 6 SQM Done", 2.1, 3.6),
                                          (2, "Floor - 4 SQM Done", 1.4, 2.4)):
            for sap, qty in ((JOB_SAPS[0][0], primer), (JOB_SAPS[1][0], top)):
                await s.execute(text(
                    'INSERT INTO consumption ("Date", "SAP_Code", "Quantity", "Site_ID", '
                    '"Tank_No", "Work_Type", "Issued_To", "Remarks") VALUES '
                    "(:d, :p, :q, :site, :t, 'Lining', 'Demo crew', :r)"),
                    {"d": day(back), "p": sap, "q": qty, "site": SITE, "t": DEMO_TAG, "r": remark})
        await s.commit()
    return 2


ROLL_SAP = "899973"


async def seed_lots(today: str) -> int:
    """Phase 16 — lots to practise FEFO on (rule 17g).

    PRACTICE PU PRIMER (899971) gets three lots, dated from TODAY so the
    buckets never drift: `PR-OLD` expired 15 days ago and still has stock (the
    red banner, and the lot FEFO must NOT suggest), `PR-SOON` expires in 20
    days (the FEFO suggestion, and the evening notice), `PR-LATE` in 9 months.
    PRACTICE CHEMOLINE (899973, ROL) has one batch of three rolls in the roll
    register — the Issue form's roll picker. One consumption names lot
    `PR-TYPO`, which no receipt brought in: the "used but never received" list.

    Ledger rows, like the other examples (a Surface Shield issue needs QC and
    an MTC this example has no business faking). Idempotent; never fatal.
    """
    import datetime as _d

    from sqlalchemy import text

    from backend.api.db import SessionLocal
    from backend.api.services import lots as LOTS
    from backend.api.services import quality

    t0 = _d.date.fromisoformat(today)

    def day(n):
        return (t0 + _d.timedelta(days=n)).isoformat()

    async with SessionLocal() as s:
        if (await s.execute(text(
                "SELECT 1 FROM lots WHERE \"Lot_Number\" = 'PR-SOON' LIMIT 1"))).first():
            return 0
        cat = await quality.controlled_category(s)
        await s.execute(text(
            'INSERT INTO inventory ("SAP_Code", "Material_Code", "Equipment_Description", '
            '"Category", "UOM", "Site_ID", "Opening_Stock", "Unit_Size", "Base_UOM") VALUES '
            "(:p, 'MAT-899973', 'PRACTICE CHEMOLINE 4MM ROLL', :c, 'ROL', :site, 0, 11, 'M2') "
            'ON CONFLICT ("SAP_Code") DO NOTHING'), {"p": ROLL_SAP, "c": cat, "site": SITE})
        await s.execute(text(
            'UPDATE inventory SET "Shelf_Life_Months" = 9 WHERE "SAP_Code" IN (\'899971\', \'899972\') '
            'AND "Shelf_Life_Months" IS NULL'))
        for lot, qty, mfd, exp in (("PR-OLD", 6, day(-290), day(-15)),
                                   ("PR-SOON", 10, day(-250), day(20)),
                                   ("PR-LATE", 12, day(-5), day(265))):
            await s.execute(text(
                'INSERT INTO receipts ("Date", "SAP_Code", "Quantity", "Site_ID", "Supplier", '
                '"Lot_Number", "Remarks") VALUES (:d, \'899971\', :q, :site, '
                "'Halcyon Industrial Supply', :l, 'Practice seed — a lot')"),
                {"d": day(-30), "q": qty, "site": SITE, "l": lot})
        await s.execute(text(
            'INSERT INTO receipts ("Date", "SAP_Code", "Quantity", "Site_ID", "Supplier", '
            '"Remarks") VALUES (:d, :p, 3, :site, \'Halcyon Industrial Supply\', '
            "'Practice seed — 3 rolls')"), {"d": day(-30), "p": ROLL_SAP, "site": SITE})
        await s.execute(text(
            'INSERT INTO consumption ("Date", "SAP_Code", "Quantity", "Site_ID", "Lot_Number", '
            '"Work_Type", "Issued_To", "Remarks") VALUES (:d, \'899971\', 1, :site, '
            "'PR-TYPO', 'Lining', 'Aria Bellweather', 'Practice seed — a lot nobody received')"),
            {"d": day(-2), "site": SITE})
        await LOTS.sync_lots_from_ledger(s, SITE)
        for lot, mfd, exp in (("PR-OLD", day(-290), day(-15)), ("PR-SOON", day(-250), day(20)),
                              ("PR-LATE", day(-5), day(265))):
            await s.execute(text(
                'UPDATE lots SET "MFD_Date" = :m, "Expiry_Date" = :e, "Expiry_Source" = \'file\' '
                'WHERE "Lot_Number" = :l AND "SAP_Code" = \'899971\' AND "Site_ID" = :site'),
                {"m": mfd, "e": exp, "l": lot, "site": SITE})
        await s.execute(text(
            'INSERT INTO lots ("Lot_Number", "SAP_Code", "Site_ID", "Received_Date", '
            '"MFD_Date", "Expiry_Date", "Expiry_Source", "Batch_Ref", "Status", "Source") '
            "VALUES ('1O26009999', :p, :site, :r, :m, :e, 'file', '1O26009999', 'open', 'lotfile') "
            'ON CONFLICT DO NOTHING'), {"p": ROLL_SAP, "site": SITE, "r": day(-30),
                                        "m": day(-200), "e": day(895)})
        for n in (1, 2, 3):
            await s.execute(text(
                'INSERT INTO lot_units ("Unit_No", "Lot_Number", "SAP_Code", "Site_ID", '
                '"Received_Date", "Location", "Source") VALUES (:u, \'1O26009999\', :p, :site, '
                ":r, 'Container 1', 'lotfile') ON CONFLICT DO NOTHING"),
                {"u": f"1O26009999{n:03d}", "p": ROLL_SAP, "site": SITE, "r": day(-30)})
        await s.commit()
    return 4


async def seed_lot_problems(today: str) -> int:
    """Phase 21a — the lot fixes, to practise on (rule 17g).

    `PR-EMPTY` (899971): 4 received, 4 issued, expires in 5 days. Before 21a it
    topped the dashboard's *Top 5 expiring lots*; now it must NOT appear there
    (nothing is left on the shelf), and it shows on the Lots page only with
    *show used-up lots*. `PR-OVER`: 2 received, three issues of 1 — the third
    draws a lot that was already used up, so the Lots page lists it under *Lot
    problems from the workbook* with its sheet and row. The Phase 16 `PR-TYPO`
    consumption gets a sheet and row too. Rows are as the Excel sync writes
    them (`Source_Sheet` / `Source_Row`). Idempotent; never fatal."""
    import datetime as _d

    from sqlalchemy import text

    from backend.api.db import SessionLocal
    from backend.api.services import lots as LOTS

    t0 = _d.date.fromisoformat(today)

    def day(n):
        return (t0 + _d.timedelta(days=n)).isoformat()

    async with SessionLocal() as s:
        if (await s.execute(text(
                "SELECT 1 FROM lots WHERE \"Lot_Number\" = 'PR-EMPTY' LIMIT 1"))).first():
            return 0
        for lot, qty in (("PR-EMPTY", 4), ("PR-OVER", 2)):
            await s.execute(text(
                'INSERT INTO receipts ("Date", "SAP_Code", "Quantity", "Site_ID", "Supplier", '
                '"Lot_Number", "Remarks") VALUES (:d, \'899971\', :q, :site, '
                "'Halcyon Industrial Supply', :l, 'Practice seed — a lot (Phase 21a)')"),
                {"d": day(-40), "q": qty, "site": SITE, "l": lot})
        draws = [("PR-EMPTY", 4, day(-10), 402), ("PR-OVER", 1, day(-9), 405),
                 ("PR-OVER", 1, day(-8), 409), ("PR-OVER", 1, day(-6), 414)]
        for lot, qty, d, row in draws:
            await s.execute(text(
                'INSERT INTO consumption ("Date", "SAP_Code", "Quantity", "Site_ID", "Lot_Number", '
                '"Work_Type", "Issued_To", "Remarks", "Source_Sheet", "Source_Row") VALUES '
                "(:d, '899971', :q, :site, :l, 'Lining', 'Aria Bellweather', "
                "'Practice seed — Phase 21a', 'Consumption Log', :r)"),
                {"d": d, "q": qty, "site": SITE, "l": lot, "r": row})
        await s.execute(text(
            'UPDATE consumption SET "Source_Sheet" = \'Consumption Log\', "Source_Row" = 418 '
            'WHERE "Lot_Number" = \'PR-TYPO\' AND "Site_ID" = :site AND "Source_Row" IS NULL'),
            {"site": SITE})
        await LOTS.sync_lots_from_ledger(s, SITE)
        for lot, exp in (("PR-EMPTY", day(5)), ("PR-OVER", day(60))):
            await s.execute(text(
                'UPDATE lots SET "MFD_Date" = :m, "Expiry_Date" = :e, "Expiry_Source" = \'file\' '
                'WHERE "Lot_Number" = :l AND "SAP_Code" = \'899971\' AND "Site_ID" = :site'),
                {"m": day(-200), "e": exp, "l": lot, "site": SITE})
        await s.commit()
    return 2


# Phase 21d — paste this into OCR Import → Paste (Consumption log) to see all
# three states (USER_MANUAL §3.17, MANUAL_TESTING_GUIDE TC-21D-*):
#     Aria, Nitril glovs, Pair, 2   → green  "learned ×3"
#     Aria, Safty goggls, Nos, 1    → gold   "did you mean Safety Goggles …?"
#     Aria, 12 inch fan, Nos, 1     → red    "not in stock — choose the item"
# (the paste lane reads Name, Product, UOM, Qty — Aria is an invented Practice name)
PRACTICE_OCR_NAMES = (("nitril glovs", "Nitril glovs", "800140", 3),
                      ("safty boots", "Safty boots", "800153", 2))


async def seed_ocr_names() -> int:
    """Phase 21d — names a store keeper 'already taught' the matcher (rule
    17g). Idempotent; never fatal."""
    from sqlalchemy import text

    from backend.api.db import SessionLocal
    n = 0
    async with SessionLocal() as s:
        for key, example, sap, times in PRACTICE_OCR_NAMES:
            if not (await s.execute(text('SELECT 1 FROM inventory WHERE "SAP_Code" = :p'),
                                    {"p": sap})).first():
                continue
            res = await s.execute(text(
                'INSERT INTO ocr_aliases ("Site_ID", written_key, written_example, "SAP_Code", '
                'confirmations, created_by, updated_by) VALUES (:s, :k, :w, :p, :n, '
                "'practice.storekeeper', 'practice.storekeeper') "
                'ON CONFLICT ("Site_ID", written_key) DO NOTHING'),
                {"s": SITE, "k": key, "w": example, "p": sap, "n": times})
            n += res.rowcount or 0
        await s.commit()
    return n


async def seed_returnables() -> int:
    """Phase 18 Track 3 — tool loans to practise the return desk on (rule 17g).

    Dated from NOW so the buckets never drift: one OVERDUE (due yesterday), one
    due back at the end of TODAY's shift, one due in three days, and one
    already returned DAMAGED (the Returned view and its condition tag). Two of
    the open loans belong to Tomas Halversen — scanning his badge (900002)
    shows a kit of two — and each carries an `Item_Ref`, so typing PR-TW-0001
    at the desk finds that loan the way a sticker scan would. Idempotent;
    never fatal.
    """
    import datetime as _d

    from sqlalchemy import text

    from backend.api.db import SessionLocal

    now = _d.datetime.now().replace(microsecond=0)
    d0 = now.replace(hour=17, minute=0, second=0)
    loans = (
        # name, Item_Ref, borrower, badge, given, due, status, condition, note
        ("PRACTICE TORQUE WRENCH", "PR-TW-0001", "Aria Bellweather", "900001",
         now - _d.timedelta(days=2), d0 - _d.timedelta(days=1), "borrowed", None, None),
        ("PRACTICE ANGLE GRINDER", "PR-AG-0002", "Tomas Halversen", "900002",
         now - _d.timedelta(hours=3), d0, "borrowed", None, None),
        ("PRACTICE SAFETY HARNESS", "PR-SH-0003", "Tomas Halversen", "900002",
         now - _d.timedelta(hours=3), d0 + _d.timedelta(days=3), "borrowed", None, None),
        ("PRACTICE IMPACT DRILL", "PR-ID-0004", "Nadia Okonjo", "900003",
         now - _d.timedelta(days=3), d0 - _d.timedelta(days=1), "returned", "damaged",
         "Chuck loose — sent for repair (Practice example)"),
    )
    async with SessionLocal() as s:
        if (await s.execute(text(
                "SELECT 1 FROM returnable_items WHERE \"Item_Ref\" = 'PR-TW-0001' LIMIT 1"))).first():
            return 0
        for name, ref, who, badge, given, due, st, cond, note in loans:
            await s.execute(text(
                'INSERT INTO returnable_items (material_name, uom, qty, borrower_name, '
                'given_time, expected_return_time, status, "Site_ID", whatsapp_alert_sent, '
                'cv_detected, cv_employee_id, "Item_Ref", returned_time, returned_by, '
                'return_condition, return_note) VALUES (:n, \'EA\', 1, :w, :g, :d, :st, :site, '
                '0, 1, :b, :r, :rt, :rb, :c, :note)'),
                {"n": name, "w": who, "g": given, "d": due, "st": st, "site": SITE,
                 "b": badge, "r": ref,
                 "rt": (due - _d.timedelta(hours=2)) if st == "returned" else None,
                 "rb": "practice.storekeeper" if st == "returned" else None,
                 "c": cond, "note": note})
        await s.commit()
    return len(loans)


async def seed_partial_return() -> int:
    """Phase 19c — a loan that came back IN PART (rule 17g).

    4 PRACTICE SCAFFOLD CLAMPS lent to Aria Bellweather (badge 900001), due 4
    days ago; 1 came back 2 days ago. The row shows **partly returned · 3 still
    out** and OVERDUE, and its slip reads "Back so far 1 of 4". It is more than
    3 days overdue, so the next morning's chase escalates it to the HOD. The
    desk returns any part of the remaining 3 (`PR-SC-0005`). Idempotent; never
    fatal."""
    import datetime as _d

    from sqlalchemy import text

    from backend.api.db import SessionLocal

    now = _d.datetime.now().replace(microsecond=0)
    async with SessionLocal() as s:
        if (await s.execute(text(
                "SELECT 1 FROM returnable_items WHERE \"Item_Ref\" = 'PR-SC-0005' LIMIT 1"))).first():
            return 0
        lid = (await s.execute(text(
            'INSERT INTO returnable_items (material_name, uom, qty, qty_returned, borrower_name, '
            'given_time, expected_return_time, status, "Site_ID", whatsapp_alert_sent, '
            'cv_detected, cv_employee_id, "Item_Ref") VALUES (\'PRACTICE SCAFFOLD CLAMPS\', '
            "'EA', 4, 1, 'Aria Bellweather', :g, :d, 'borrowed', :site, 1, 1, '900001', "
            "'PR-SC-0005') RETURNING id"),
            {"g": now - _d.timedelta(days=6), "d": now - _d.timedelta(days=4),
             "site": SITE})).scalar_one()
        await s.execute(text(
            'INSERT INTO returnable_returns (loan_id, qty, condition, note, returned_by, '
            'returned_time, "Site_ID") VALUES (:l, 1, \'ok\', \'Practice example — 1 of 4 back\', '
            "'practice.storekeeper', :t, :site)"),
            {"l": lid, "t": now - _d.timedelta(days=2), "site": SITE})
        await s.commit()
    return 1


REORDER_SAPS = ("899981", "899982", "899983")    # synthetic, Practice-only (P12-5)


async def seed_reorder_examples() -> int:
    """Phase 18 Track 4 — one item per reorder colour (rule 17g).

    The tutorial dataset's consumption has FIXED dates (P12-5), so as the
    90-day window moves on, its items drift to "no recent use". These three
    are dated from TODAY so Stock → Reorder signals always shows one of each:
    each used 3 a day for the last 30 days (minimum 90 at 30 days of cover),
    holding 20 (red), 120 (amber) and 300 (green). Idempotent; never fatal.
    """
    import datetime as _d

    from sqlalchemy import text

    from backend.api.db import SessionLocal

    t0 = _d.date.today()
    async with SessionLocal() as s:
        if (await s.execute(text(
                'SELECT 1 FROM inventory WHERE "SAP_Code" = :p'), {"p": REORDER_SAPS[0]})).first():
            return 0
        for sap, name, hold in zip(REORDER_SAPS, ("PRACTICE CABLE TIES (red)",
                                                  "PRACTICE MASKING TAPE (amber)",
                                                  "PRACTICE NITRILE GLOVES (green)"), (20, 120, 300)):
            await s.execute(text(
                'INSERT INTO inventory ("SAP_Code", "Material_Code", "Equipment_Description", '
                '"Category", "UOM", "Site_ID", "Opening_Stock") VALUES '
                "(:p, :m, :n, 'R/L Consumables', 'EA', :site, 0) ON CONFLICT DO NOTHING"),
                {"p": sap, "m": f"MAT-{sap}", "n": name, "site": SITE})
            await s.execute(text(
                'INSERT INTO receipts ("Date", "SAP_Code", "Quantity", "Site_ID", "Supplier", '
                '"Remarks") VALUES (:d, :p, :q, :site, \'Halcyon Industrial Supply\', '
                "'Practice seed — reorder example')"),
                {"d": (t0 - _d.timedelta(days=35)).isoformat(), "p": sap, "q": hold + 90, "site": SITE})
            for k in range(30):
                await s.execute(text(
                    'INSERT INTO consumption ("Date", "SAP_Code", "Quantity", "Site_ID", '
                    '"Work_Type", "Issued_To", "Remarks") VALUES (:d, :p, 3, :site, '
                    "'Maintenance', 'Aria Bellweather', 'Practice seed — reorder example')"),
                    {"d": (t0 - _d.timedelta(days=k + 1)).isoformat(), "p": sap, "site": SITE})
        await s.commit()
    return len(REORDER_SAPS)


ACCEPTED_SAP = "899984"                          # synthetic, Practice-only (P12-5)


async def seed_accepted_minimum() -> int:
    """Phase 19a — an HOD-accepted minimum the system now flags CHANGED
    (rule 17g).

    PRACTICE SAFETY GLASSES are used 3 a day (recommended 90). practice.hod
    accepted 60 for CNCEC weeks ago. 90 is 50 % above 60, beyond the ±20 %
    band, so the row shows **accepted** + **changed**, and with 80 in stock
    against 60 it is amber. Idempotent; never fatal."""
    import datetime as _d

    from sqlalchemy import text

    from backend.api.db import SessionLocal

    t0 = _d.date.today()
    async with SessionLocal() as s:
        if (await s.execute(text(
                'SELECT 1 FROM inventory WHERE "SAP_Code" = :p'), {"p": ACCEPTED_SAP})).first():
            return 0
        await s.execute(text(
            'INSERT INTO inventory ("SAP_Code", "Material_Code", "Equipment_Description", '
            '"Category", "UOM", "Site_ID", "Opening_Stock") VALUES '
            "(:p, :m, 'PRACTICE SAFETY GLASSES (accepted, changed)', 'R/L Consumables', 'EA', "
            ":site, 0) ON CONFLICT DO NOTHING"),
            {"p": ACCEPTED_SAP, "m": f"MAT-{ACCEPTED_SAP}", "site": SITE})
        await s.execute(text(
            'INSERT INTO receipts ("Date", "SAP_Code", "Quantity", "Site_ID", "Supplier", '
            '"Remarks") VALUES (:d, :p, 170, :site, \'Halcyon Industrial Supply\', '
            "'Practice seed — accepted minimum')"),
            {"d": (t0 - _d.timedelta(days=35)).isoformat(), "p": ACCEPTED_SAP, "site": SITE})
        for k in range(30):
            await s.execute(text(
                'INSERT INTO consumption ("Date", "SAP_Code", "Quantity", "Site_ID", '
                '"Work_Type", "Issued_To", "Remarks") VALUES (:d, :p, 3, :site, '
                "'Maintenance', 'Aria Bellweather', 'Practice seed — accepted minimum')"),
                {"d": (t0 - _d.timedelta(days=k + 1)).isoformat(), "p": ACCEPTED_SAP, "site": SITE})
        await s.execute(text(
            'INSERT INTO inventory_site_overrides ("SAP_Code", "Site_ID", "Minimum_Qty", '
            "updated_by, updated_at) VALUES (:p, :site, 60, 'practice.hod', :at) "
            'ON CONFLICT ("SAP_Code", "Site_ID") DO NOTHING'),
            {"p": ACCEPTED_SAP, "site": SITE,
             "at": _d.datetime.combine(t0 - _d.timedelta(days=40), _d.time(9, 0))})
        await s.commit()
    return 1


async def seed_on_order_examples() -> int:
    """Phase 19b — on order per site vs GLOBAL (rule 17g).

    PRACTICE MASKING TAPE (amber) has 25 open on a PO raised from a CNCEC PR:
    its suggested order drops from 60 to 35. PRACTICE CABLE TIES (red) has 40
    open on a PO with NO PR: shown as "+ 40 global", while its suggested order
    stays 160. Idempotent; never fatal."""
    from sqlalchemy import text

    from backend.api.db import SessionLocal

    async with SessionLocal() as s:
        if (await s.execute(text(
                "SELECT 1 FROM purchase_orders WHERE \"PO_Number\" = 'PRACTICE-PO-19B-1'"))).first():
            return 0
        await s.execute(text(
            'INSERT INTO pr_registry ("PR_Number", "Site_ID", created_by) VALUES '
            "('PRACTICE-PR-19B', :site, 'practice.hod') ON CONFLICT DO NOTHING"), {"site": SITE})
        await s.execute(text(
            'INSERT INTO pr_master ("PR_Number", "SAP_Code", "Requested_Qty", "Site_ID", '
            '"Material_Code", "Material_Name", status) VALUES (\'PRACTICE-PR-19B\', :p, 25, '
            ":site, :m, 'PRACTICE MASKING TAPE (amber)', 'approved')"),
            {"p": REORDER_SAPS[1], "m": f"MAT-{REORDER_SAPS[1]}", "site": SITE})
        await s.execute(text(
            'INSERT INTO purchase_orders ("PO_Number", "PR_Number", "Site_ID", "Vendor_Name") VALUES '
            "('PRACTICE-PO-19B-1', 'PRACTICE-PR-19B', :site, 'Halcyon Industrial Supply'), "
            "('PRACTICE-PO-19B-2', NULL, NULL, 'Halcyon Industrial Supply')"), {"site": SITE})
        await s.execute(text(
            'INSERT INTO po_items ("PO_Number", line_no, "Material_Code", "Description", "Qty", '
            '"UOM", "PR_Number", line_status) VALUES '
            "('PRACTICE-PO-19B-1', 1, :m1, 'PRACTICE MASKING TAPE (amber)', 25, 'EA', "
            "'PRACTICE-PR-19B', 'open'), "
            "('PRACTICE-PO-19B-2', 1, :m0, 'PRACTICE CABLE TIES (red)', 40, 'EA', NULL, 'open')"),
            {"m0": f"MAT-{REORDER_SAPS[0]}", "m1": f"MAT-{REORDER_SAPS[1]}"})
        await s.commit()
    return 2


PRACTICE_DN = "90001"


def _practice_png(lines: list[str]) -> bytes:
    """A drawn, obviously synthetic document (no real company — P12-0)."""
    import io

    from PIL import Image, ImageDraw
    img = Image.new("RGB", (900, 600), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([20, 20, 880, 580], outline="black", width=3)
    y = 50
    for ln in lines:
        d.text((50, y), ln, fill="black")
        y += 40
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _practice_pdf(lines: list[str]) -> bytes:
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for ln in lines:
        pdf.cell(0, 9, ln, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _practice_xlsx(rows: list[list]) -> bytes:
    import io

    import openpyxl
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def seed_drive_examples(today: str) -> dict:
    """Phase 22 (rule 17g) — the Drive features, on Practice data.

    Practice never pulls from Drive (rule 17); these are what a pull WOULD
    have indexed, drawn here, all synthetic:
      · DN — receipt DN 90001 with a drawn delivery note to open (📎), and a
        delivery with NO DN that gets its WD number (22b);
      · MTC — a certificate whose batch and product match lot PR-SOON (filed
        by itself), and one with no batch ("container 1") proposed by the HOD
        on PR-LATE for the QC trainee to confirm (22c);
      · a request workbook from 10 days ago: a general item without a PR, one
        line with no SAP code (22d);
      · the site's Day / Night preparers, a learned tank spelling, and a NIGHT
        paper two days ago already in the "workbook" — paste it to see Compare
        (22e). Idempotent; never fatal.
    """
    import datetime as _d

    from sqlalchemy import text

    from backend.api.db import SessionLocal
    from backend.api.services import drive_links, drive_sync as DS
    from backend.api.services import ledger, mtc_links, preparers, requests_sync

    t0 = _d.date.fromisoformat(today)
    out = {}
    folder = DS.CACHE_DIR / "practice"
    folder.mkdir(parents=True, exist_ok=True)

    async def put_file(s, drive_id, kind, name, data, mime):
        path = folder / f"{drive_id}{pathlib.Path(name).suffix}"
        path.write_bytes(data)
        await s.execute(text('''
            INSERT INTO drive_files (drive_id, kind, folder, name, mime, size, cache_path)
            VALUES (:d, :k, 'Practice', :n, :m, :z, :p)
            ON CONFLICT (drive_id) DO UPDATE SET cache_path = :p, removed_at = NULL'''),
            {"d": drive_id, "k": kind, "n": name, "m": mime, "z": len(data), "p": str(path)})

    async with SessionLocal() as s:
        # ── 22b: a DN to open, and a delivery without one ───────────────────
        if not (await s.execute(text('SELECT 1 FROM receipts WHERE "DN_No" = :d'),
                                {"d": PRACTICE_DN})).first():
            for dn, veh, qty in ((PRACTICE_DN, "PR-1001", 5.0), ("WD", "PR-2002", 3.0)):
                await ledger.post_receipt(s, username="practice.storekeeper", data={
                    "Date": (t0 - _d.timedelta(days=3)).isoformat(), "SAP_Code": "899001",
                    "Quantity": qty, "Site_ID": SITE, "Supplier": "Halcyon Industrial Supply",
                    "Remarks": "Practice — DN example",
                    "extra": {"DN_No": dn, "Vehicle_No": veh}})
        await put_file(s, "practice-dn-90001", "dn", f"DN# {PRACTICE_DN}.png",
                       _practice_png(["DELIVERY NOTE", f"DN# {PRACTICE_DN}", "PRACTICE DATA - NOT A REAL DOCUMENT",
                                      "Tutorial Demo Gasket Set (899001)   x 5"]), "image/png")
        # ── 22c: two certificates ──────────────────────────────────────────
        lot = (await s.execute(text(
            'SELECT "Expiry_Date" FROM lots WHERE "Lot_Number" = \'PR-SOON\' LIMIT 1'))).scalar()
        if lot:
            exp = _d.date.fromisoformat(str(lot)[:10])
            await put_file(s, "practice-mtc-1", "mtc", "PRACTICE PU PRIMER (BNO-PR-SOON).pdf",
                           _practice_pdf(["TEST CERTIFICATE - PRACTICE DATA",
                                          "Product Name : PRACTICE PU PRIMER", "Batch No : PR-SOON",
                                          f"D.O.M : {(exp - _d.timedelta(days=270)).strftime('%d/%m/%Y')}",
                                          f"D.O.E : {exp.strftime('%d/%m/%Y')}"]), "application/pdf")
        await put_file(s, "practice-mtc-2", "mtc", "PRACTICE MTC - container 1.pdf",
                       _practice_pdf(["TEST CERTIFICATE - PRACTICE DATA", "Container 1",
                                      "(no batch number on this page)"]), "application/pdf")
        # ── 22d: a request mailed 10 days ago ──────────────────────────────
        codes = [r[0] for r in (await s.execute(text(
            'SELECT "Material_Code" FROM inventory WHERE "Site_ID" = :s AND "Material_Code" IS NOT NULL '
            "AND \"Category\" NOT ILIKE '%surface%' ORDER BY \"SAP_Code\" LIMIT 2"), {"s": SITE})).all()]
        req_day = t0 - _d.timedelta(days=10)
        rows = [["Sl. No.", "Material Code", "Material Description", "UOM", "Qty", "Pending Qty", "Remarks"]]
        for i, c in enumerate(codes, 1):
            rows.append([i, c, f"Practice requested item {i}", "EA", 20 * i, 20 * i, None])
        rows.append([len(rows), "N/A", "Practice head pan (no SAP code yet)", "Nos", 10, 10, None])
        await put_file(s, "practice-req-1", "pending", f"Request {req_day.strftime('%d-%m-%Y')}.xlsx",
                       _practice_xlsx(rows),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        await s.commit()
        out.update(await drive_links.relink_all(s))
        await s.commit()
        # QC's queue: the no-batch certificate proposed on PR-LATE by the HOD
        fid = (await s.execute(text("SELECT id FROM drive_files WHERE drive_id = 'practice-mtc-2'"))).scalar()
        late = (await s.execute(text(
            'SELECT "SAP_Code", COALESCE("Site_ID", \'HQ\') FROM lots WHERE "Lot_Number" = \'PR-LATE\' LIMIT 1'))).first()
        if fid and late:
            await s.execute(text('''
                INSERT INTO mtc_assignments (drive_file_id, "SAP_Code", "Lot_Number", "Site_ID",
                    source, status, proposed_by, note)
                SELECT :f, :s, 'PR-LATE', :site, 'manual', 'proposed', 'practice.hod', 'container 1'
                WHERE NOT EXISTS (SELECT 1 FROM mtc_assignments WHERE drive_file_id = :f)'''),
                {"f": fid, "s": late[0], "site": late[1]})
        # ── 22e: the preparers, a learned tank, a Night paper in the workbook ──
        if not await preparers.history(s, SITE):
            await preparers.save(s, SITE, [{"from": "", "day": "Practice Day", "night": "Practice Night"}])
        tag = (await s.execute(text(
            'SELECT "Equipment_Tag_No" FROM sme_equipment WHERE "Site_ID" = :s ORDER BY 1 LIMIT 1'),
            {"s": SITE})).scalar()
        if tag:
            from backend.api.ai.paper_fields import tank_key
            await s.execute(text('''
                INSERT INTO sme_tank_alias ("Site_ID", alias_raw, alias_norm, "Equipment_Tag_No", status,
                                            resolved_by, resolved_at)
                VALUES (:s, :raw, :n, :t, 'mapped', 'practice.storekeeper', CURRENT_TIMESTAMP)
                ON CONFLICT ("Site_ID", alias_norm) DO NOTHING'''),
                {"s": SITE, "raw": "Tank one", "n": tank_key("Tank one"), "t": tag})
        night = (t0 - _d.timedelta(days=2)).isoformat()
        gl = (await s.execute(text(
            'SELECT "SAP_Code" FROM inventory WHERE "Site_ID" = :s AND "Category" NOT ILIKE \'%surface%\' '
            'ORDER BY "SAP_Code" LIMIT 2'), {"s": SITE})).all()
        if not (await s.execute(text(
                "SELECT 1 FROM consumption WHERE \"Source_Ref\" LIKE 'XLSX:PRACTICE:night:%' LIMIT 1"))).first():
            for n, (sap,) in enumerate(gl, start=1):
                await s.execute(text('''
                    INSERT INTO consumption ("Date", "SAP_Code", "Quantity", "Site_ID", "Work_Type",
                        "Issued_To", "Issued_By", "Prepared_By", "Tank_No", "Source_Ref",
                        "Source_Sheet", "Source_Row", "Item_Type")
                    VALUES (:d, :sap, :q, :s, 'R/L', 'Aria Bellweather', 'Practice Night', 'Practice Night',
                            :tank, :ref, 'Consumption Log', :row, 'R/L Consumables')'''),
                    {"d": night, "sap": sap, "q": float(n), "s": SITE, "tank": tag,
                     "ref": f"XLSX:PRACTICE:night:{night}:{sap}:{n}", "row": 900 + n})
        await s.commit()
        out["night_paper"] = night
    _ = (mtc_links, requests_sync)          # imported for their linkers
    return out


def fixture_version() -> int:
    sys.path.insert(0, str(_ROOT / "tools"))
    import make_tutorial_db
    return make_tutorial_db.DATASET_VERSION


async def run() -> int:
    from backend.api.config import database_name, is_practice
    if not is_practice():
        print("❌ practice_overlay runs as a Practice process: set GI_INSTANCE=training",
              file=sys.stderr)
        return 2
    from backend.api.db import SessionLocal, engine
    print(f"▶ overlay v{OVERLAY_VERSION} → {database_name()}")
    async with SessionLocal() as s:
        acc = await seed_accounts(s)
        await seed_settings(s, fixture_version=fixture_version())
        await s.commit()
    counts = await seed_queues()
    await engine.dispose()
    print(f"  accounts={acc['accounts']} · " + " · ".join(f"{k}={v}" for k, v in counts.items()))
    if acc["admin_password"]:
        print(f"  🔑 practice.admin password (generated — shown ONCE): {acc['admin_password']}")
    # ⚠️ THE OVERLAY CHECKS ITSELF, as the tutorial fixture does: an empty
    # approval queue does not fail anything downstream — it produces a Practice
    # environment where an HOD trainee has nothing to practise on, which looks
    # finished and is not.
    problems = []
    if acc["accounts"] != len(ACCOUNTS):
        problems.append(f"expected {len(ACCOUNTS)} practice accounts, found {acc['accounts']}")
    for q in ("pending_receipts", "pending_issues", "pending_returns", "pr_draft"):
        if not counts.get(q):
            problems.append(f"{q} is EMPTY — that role's queue would have nothing to practise on")
    if problems:
        print("❌ overlay self-check failed:\n  " + "\n  ".join(problems), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(asyncio.run(run()))
