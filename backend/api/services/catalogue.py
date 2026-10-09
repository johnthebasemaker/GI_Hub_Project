"""
backend/api/services/catalogue.py — the material catalogue, the site equipment
list and their pictures, from Drive (Phase 23d, rulings Q23-5..9).

THE CATALOGUE. Drive's root holds `All MATERIAL CODES-15.04.2026.xlsx`: two
sheets, `7-SERIES` (5,976 rows) and `6-SERIES` (2,415) — and the 7-SERIES sheet
already contains every 6-SERIES row, so the file is 5,976 unique codes. Columns
`Material · Material Description · UOM`. The NEWEST edition wins (by the date
in its name, else Drive's modified time), so a new file needs no change here.
The two sheets are merged; a code whose descriptions disagree is REPORTED, the
7-SERIES spelling kept. A code that disappears from a newer edition is marked
`removed_at`, never deleted (a request may still name it).

THE EQUIPMENT LIST (Q23-9). `Equipment list Updated as on 06-09-2026.xlsx`: the
site's plant and tools under section headings (Vehicle, Blasting Equipment,
Utilities …). S.No repeats in the sheet, so a line is keyed by its section and
its description. Not the SME tanks (`Equipment.xlsx`).

PICTURES FROM DRIVE (Q23-6). A file in Drive's `Material Images` folder whose
name STARTS with a GI code (`GI-7000003.jpg`, `GI-7000003 tyvek front.png`)
becomes that code's picture on the next pull. Never a web search.
"""
from __future__ import annotations

import datetime as _dt
import io
import re
import warnings
from pathlib import Path
from typing import Any, Optional

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

SITE = "CNCEC"
CODE_RX = re.compile(r"^\s*(GI-\d{5,})", re.I)
CATALOGUE_NAME = re.compile(r"^all\s*material\s*codes", re.I)
EQUIPMENT_NAME = re.compile(r"^equipment\s*list", re.I)
MAX_PER_ITEM = 4


def _s(v) -> str:
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def name_date(name: str) -> Optional[_dt.date]:
    """`…-15.04.2026.xlsx` / `… as on 06-09-2026.xlsx` → that date (day first)."""
    m = re.search(r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})", name or "")
    if not m:
        return None
    try:
        return _dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def newest(files: list[dict]) -> Optional[dict]:
    """The edition to read: the latest date in the name, then Drive's time."""
    if not files:
        return None
    return max(files, key=lambda f: (name_date(f["name"]) or _dt.date.min, f.get("modified_time") or ""))


def read_catalogue(data: bytes) -> dict:
    """{codes: {code: {description, uom, series}}, conflicts: [...], rows: n}"""
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    codes: dict[str, dict] = {}
    conflicts: list[dict] = []
    rows = 0
    # 7-SERIES first: it is the superset, so its spelling is the one kept
    sheets = sorted(wb.worksheets, key=lambda w: (not w.title.upper().startswith("7"), w.title))
    for ws in sheets:
        header = None
        for r in ws.iter_rows(values_only=True):
            cells = [_s(c) for c in r[:6]]
            if header is None:
                low = [c.lower() for c in cells]
                if "material" in low and any("description" in c for c in low):
                    header = {"code": low.index("material"),
                              "desc": next(i for i, c in enumerate(low) if "description" in c),
                              "uom": low.index("uom") if "uom" in low else None}
                continue
            code = cells[header["code"]].upper() if header["code"] < len(cells) else ""
            if not CODE_RX.match(code):
                continue
            rows += 1
            desc = cells[header["desc"]]
            uom = cells[header["uom"]] if header["uom"] is not None else ""
            if code in codes:
                if codes[code]["description"].upper() != desc.upper() and desc:
                    conflicts.append({"code": code, "kept": codes[code]["description"],
                                      "other": desc, "sheet": ws.title})
                continue
            codes[code] = {"description": desc, "uom": uom, "series": code[3:4]}
    return {"codes": codes, "conflicts": conflicts, "rows": rows}


def _equipment_key(category: str, description: str) -> str:
    return f"{_s(category).lower()}|{re.sub(r'[^a-z0-9]+', ' ', description.lower()).strip()}"


def read_equipment(data: bytes) -> dict:
    """{items: [...], site: …} from the plant & tools list (Q23-9)."""
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    items: list[dict] = []
    head = None
    category = ""
    site = SITE
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        cells = list(r[:14]) + [None] * max(0, 14 - len(r))
        txt = [_s(c) for c in cells]
        joined = " ".join(t for t in txt if t)
        if head is None:
            m = re.search(r"AT\s+([A-Z0-9]+)\s+SITE", joined.upper())
            if m:
                site = m.group(1)
            low = [t.lower() for t in txt]
            if any(t.startswith("sl") for t in low) and any("description" in t for t in low):
                head = {"sl": next(k for k, t in enumerate(low) if t.startswith("sl")),
                        "desc": next(k for k, t in enumerate(low) if "description" in t)}
                for name, words in (("brand", ("brand",)), ("serial", ("serial",)),
                                    ("asset", ("asset",)), ("uom", ("uom",)), ("qty", ("qty",)),
                                    ("sticker", ("sticker",)), ("condition", ("condition",)),
                                    ("remarks", ("remark",))):
                    head[name] = next((k for k, t in enumerate(low) if any(w in t for w in words)), None)
                head["expiry"] = head["sticker"] + 1 if head.get("sticker") is not None else None
            continue
        desc = txt[head["desc"]]
        sl = txt[head["sl"]]
        if not desc and sl and not re.match(r"^\d+$", sl):
            category = sl                       # a section heading: "Vehicle", "Utilities" …
            continue
        if not desc or desc.lower() in ("sticker no", "expired date"):
            continue

        def col(name):
            k = head.get(name)
            v = cells[k] if k is not None else None
            if isinstance(v, (_dt.date, _dt.datetime)):
                return v.strftime("%Y-%m-%d")
            return _s(v) or None
        qty = cells[head["qty"]] if head.get("qty") is not None else None
        try:
            qty = float(qty) if qty not in (None, "") else None
        except (TypeError, ValueError):
            qty = None
        items.append({"key": _equipment_key(category, desc), "category": category or None,
                      "description": desc, "brand": col("brand"), "serials": col("serial"),
                      "asset_no": col("asset"), "uom": col("uom"), "qty": qty,
                      "sticker_no": col("sticker"), "sticker_expiry": col("expiry"),
                      "condition": col("condition"), "remarks": col("remarks"), "row": i})
    # the same description twice in one section stays two lines, keyed apart
    seen: dict[str, int] = {}
    for it in items:
        n = seen.get(it["key"], 0)
        seen[it["key"]] = n + 1
        if n:
            it["key"] = f"{it['key']}#{n + 1}"
    return {"items": items, "site": site}


# ── linkers (run after every Drive pull) ─────────────────────────────────────
def _files(rows) -> list[dict]:
    return [dict(r) for r in rows]


async def link_catalogue(session: AsyncSession, *, read_bytes=None) -> dict:
    read_bytes = read_bytes or (lambda p: Path(p).read_bytes())
    files = _files((await session.execute(text(
        "SELECT id, name, cache_path, modified_time FROM drive_files WHERE kind = 'catalogue' "
        "AND removed_at IS NULL AND cache_path IS NOT NULL"))).mappings().all())
    f = newest([x for x in files if CATALOGUE_NAME.search(x["name"])])
    if not f:
        return {"file": None}
    try:
        parsed = read_catalogue(read_bytes(f["cache_path"]))
    except Exception as e:  # noqa: BLE001
        return {"file": f["name"], "error": f"unreadable: {type(e).__name__}"}
    if not parsed["codes"]:
        return {"file": f["name"], "error": "no GI codes found"}
    for code, v in parsed["codes"].items():
        await session.execute(text('''
            INSERT INTO material_catalog ("Material_Code", description, uom, series, source_file)
            VALUES (:c, :d, :u, :s, :f)
            ON CONFLICT ("Material_Code") DO UPDATE SET description = :d, uom = :u, series = :s,
                source_file = :f, last_seen = CURRENT_TIMESTAMP, removed_at = NULL'''),
            {"c": code, "d": v["description"], "u": v["uom"], "s": v["series"], "f": f["name"]})
    gone = (await session.execute(text(
        'UPDATE material_catalog SET removed_at = CURRENT_TIMESTAMP WHERE removed_at IS NULL '
        'AND NOT ("Material_Code" = ANY(:codes)) RETURNING "Material_Code"'),
        {"codes": list(parsed["codes"])})).all()
    await session.execute(text(
        "UPDATE drive_files SET link_status = CASE WHEN id = :i THEN 'linked' ELSE 'other' END "
        "WHERE kind = 'catalogue'"), {"i": f["id"]})
    not_in = [r[0] for r in (await session.execute(text(
        'SELECT DISTINCT UPPER(TRIM("Material_Code")) FROM inventory WHERE COALESCE("Material_Code", \'\') '
        '~* \'^GI-[0-9]+\' AND UPPER(TRIM("Material_Code")) NOT IN '
        '(SELECT "Material_Code" FROM material_catalog WHERE removed_at IS NULL)'))).all()]
    rep = {"file": f["name"], "codes": len(parsed["codes"]), "rows": parsed["rows"],
           "conflicts": parsed["conflicts"][:50], "conflict_count": len(parsed["conflicts"]),
           "removed": len(gone), "item_master_not_in_catalogue": sorted(not_in),
           "at": _dt.datetime.now().isoformat(timespec="seconds")}
    await _save_report(session, "catalogue", rep)
    return rep


REPORT_KEY = "catalogue_report:{part}"


async def _save_report(session: AsyncSession, part: str, rep: dict) -> None:
    import json
    await session.execute(text(
        "INSERT INTO app_settings (key, value) VALUES (:k, :v) ON CONFLICT (key) DO UPDATE SET value = :v"),
        {"k": REPORT_KEY.format(part=part), "v": json.dumps(rep, default=str)})


async def report(session: AsyncSession) -> dict:
    import json
    out = {}
    for part in ("catalogue", "equipment"):
        raw = (await session.execute(text("SELECT value FROM app_settings WHERE key = :k"),
                                     {"k": REPORT_KEY.format(part=part)})).scalar()
        try:
            out[part] = json.loads(raw) if raw else None
        except ValueError:
            out[part] = None
    return out


async def link_equipment(session: AsyncSession, *, read_bytes=None) -> dict:
    read_bytes = read_bytes or (lambda p: Path(p).read_bytes())
    files = _files((await session.execute(text(
        "SELECT id, name, cache_path, modified_time FROM drive_files WHERE kind = 'equipment' "
        "AND removed_at IS NULL AND cache_path IS NOT NULL"))).mappings().all())
    f = newest([x for x in files if EQUIPMENT_NAME.search(x["name"])])
    if not f:
        return {"file": None}
    try:
        parsed = read_equipment(read_bytes(f["cache_path"]))
    except Exception as e:  # noqa: BLE001
        return {"file": f["name"], "error": f"unreadable: {type(e).__name__}"}
    site = parsed["site"]
    for it in parsed["items"]:
        await session.execute(text('''
            INSERT INTO site_equipment ("Site_ID", equipment_key, category, description, brand,
                serials, asset_no, uom, qty, sticker_no, sticker_expiry, condition, remarks,
                source_file, source_row)
            VALUES (:s, :k, :cat, :d, :b, :ser, :a, :u, :q, :st, :ex, :co, :rm, :f, :row)
            ON CONFLICT ("Site_ID", equipment_key) DO UPDATE SET category = :cat, description = :d,
                brand = :b, serials = :ser, asset_no = :a, uom = :u, qty = :q, sticker_no = :st,
                sticker_expiry = :ex, condition = :co, remarks = :rm, source_file = :f,
                source_row = :row, last_seen = CURRENT_TIMESTAMP, removed_at = NULL'''),
            {"s": site, "k": it["key"], "cat": it["category"], "d": it["description"],
             "b": it["brand"], "ser": it["serials"], "a": it["asset_no"], "u": it["uom"],
             "q": it["qty"], "st": it["sticker_no"], "ex": it["sticker_expiry"],
             "co": it["condition"], "rm": it["remarks"], "f": f["name"], "row": it["row"]})
    await session.execute(text(
        'UPDATE site_equipment SET removed_at = CURRENT_TIMESTAMP WHERE "Site_ID" = :s '
        'AND removed_at IS NULL AND NOT (equipment_key = ANY(:keys))'),
        {"s": site, "keys": [it["key"] for it in parsed["items"]] or [""]})
    await session.execute(text(
        "UPDATE drive_files SET link_status = CASE WHEN id = :i THEN 'linked' ELSE 'other' END "
        "WHERE kind = 'equipment'"), {"i": f["id"]})
    rep = {"file": f["name"], "site": site, "items": len(parsed["items"]),
           "at": _dt.datetime.now().isoformat(timespec="seconds")}
    await _save_report(session, "equipment", rep)
    return rep


async def link_images(session: AsyncSession, *, read_bytes=None) -> dict:
    """Drive's `Material Images` folder: a file named after a GI code becomes
    that code's picture; a file that leaves Drive takes its picture with it."""
    from . import media
    read_bytes = read_bytes or (lambda p: Path(p).read_bytes())
    files = _files((await session.execute(text(
        "SELECT id, name, cache_path, removed_at FROM drive_files WHERE kind = 'images' "
        "AND cache_path IS NOT NULL"))).mappings().all())
    have = {r[0] for r in (await session.execute(text(
        "SELECT drive_file_id FROM item_images WHERE source = 'drive' AND drive_file_id IS NOT NULL "
        "AND removed_at IS NULL"))).all()}
    out = {"files": len(files), "added": 0, "unnamed": [], "unreadable": [], "removed": 0}
    for f in files:
        if f["removed_at"]:
            n = (await session.execute(text(
                "UPDATE item_images SET removed_at = CURRENT_TIMESTAMP, removed_by = 'drive' "
                "WHERE source = 'drive' AND drive_file_id = :i AND removed_at IS NULL"),
                {"i": f["id"]})).rowcount
            out["removed"] += n or 0
            continue
        if f["id"] in have:
            continue
        m = CODE_RX.match(f["name"])
        if not m:
            out["unnamed"].append(f["name"])
            await session.execute(text("UPDATE drive_files SET link_status = 'unlinked' WHERE id = :i"),
                                  {"i": f["id"]})
            continue
        try:
            st = media.store(read_bytes(f["cache_path"]))
        except (media.MediaError, OSError):
            out["unreadable"].append(f["name"])
            continue
        await add_image(session, "material", m.group(1).upper(), st, source="drive",
                        user="drive", drive_file_id=f["id"])
        await session.execute(text("UPDATE drive_files SET link_status = 'linked' WHERE id = :i"),
                              {"i": f["id"]})
        out["added"] += 1
    return out


# ── pictures ─────────────────────────────────────────────────────────────────
async def live_images(session: AsyncSession, kind: str, key: str) -> list[dict]:
    return [dict(r) for r in (await session.execute(text(
        "SELECT * FROM item_images WHERE kind = :k AND item_key = :i AND removed_at IS NULL "
        "ORDER BY is_primary DESC, uploaded_at, id"), {"k": kind, "i": key})).mappings().all()]


async def add_image(session: AsyncSession, kind: str, key: str, st, *, source: str, user: str,
                    drive_file_id: Optional[int] = None, caption: Optional[str] = None) -> dict:
    """One more picture for an item (the same picture twice is not added
    twice). The first picture becomes the primary. ValueError past 4."""
    live = await live_images(session, kind, key)
    dup = next((x for x in live if x["sha256"] == st.sha256), None)
    if dup:
        return dup
    if len(live) >= MAX_PER_ITEM:
        raise ValueError(f"an item has at most {MAX_PER_ITEM} pictures — remove one first")
    row = (await session.execute(text('''
        INSERT INTO item_images (kind, item_key, sha256, mime, width, height, bytes, source,
                                 drive_file_id, is_primary, caption, uploaded_by)
        VALUES (:k, :i, :sha, :m, :w, :h, :b, :src, :df, :p, :c, :u) RETURNING *'''),
        {"k": kind, "i": key, "sha": st.sha256, "m": st.mime, "w": st.width, "h": st.height,
         "b": st.bytes, "src": source, "df": drive_file_id, "p": not live, "c": caption,
         "u": user})).mappings().first()
    return dict(row)


def stem(description: str) -> str:
    """A family's shared name: sizes, thicknesses and grades taken off the end
    (`RUBBER SHEET VE611BN-5MM` → `rubber sheet ve611bn`,
    `SAFETY COVERALL XL` → `safety coverall`)."""
    d = (description or "").lower()
    d = re.sub(r"[-\s]*\d+(\.\d+)?\s*(mm|cm|m|in|inch|ltr?s?|l|kg|gr?|g|mtrs?|ml)\b.*$", "", d)
    d = re.sub(r"\b(xxxl|xxl|xl|xs|s|m|l|small|medium|large|size\s*\w+)\b\s*$", "", d)
    d = re.sub(r"[^a-z0-9]+", " ", d).strip()
    return d


async def family_suggestions(session: AsyncSession, code: str, limit: int = 6) -> list[dict]:
    """Codes of the same family that HAVE a picture — offered, never applied."""
    me = (await session.execute(text(
        'SELECT description FROM material_catalog WHERE "Material_Code" = :c'), {"c": code})).scalar()
    if not me:
        return []
    st = stem(me)
    if len(st) < 5:
        return []
    rows = (await session.execute(text('''
        SELECT c."Material_Code" AS code, c.description, i.id AS image_id
        FROM material_catalog c JOIN item_images i
          ON i.kind = 'material' AND i.item_key = c."Material_Code" AND i.removed_at IS NULL AND i.is_primary
        WHERE c."Material_Code" <> :c AND c.removed_at IS NULL
          AND lower(c.description) LIKE :p
        ORDER BY c."Material_Code" LIMIT :n'''), {"c": code, "p": st.split(" ")[0] + "%", "n": 50})).mappings().all()
    return [dict(r) for r in rows if stem(r["description"]) == st][:limit]


def as_row(r: Any) -> dict:
    return dict(r)
