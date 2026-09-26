"""
backend/api/services/stock_excel.py — why GI Hub's stock and the Excel
workbook's "Current Stock" disagree, SAP by SAP, in words someone can act on.

The Excel sync ends with a stock verification (`tools/pg_excel_sync.py`): for
every SAP, the workbook's Inventory sheet "Current Stock" against GI Hub's
`Opening_Stock + Σreceipts − Σconsumption − Σreturns`. That check only said
WHICH SAPs differ. This module says WHY, by decomposing each difference into
signed causes whose effects add up to it:

  app_only      a movement entered in GI Hub that the workbook's log never got
                (receipts via a DN, a return, an issue, a test entry …)
  vanished      a line GI Hub took from the workbook that the workbook no
                longer has (a deleted or re-dated row in the log)
  not_synced    a workbook log line GI Hub does not have yet — run the sync
  qty_edited    a log line whose quantity was edited since the last sync
  sheet_totals  the Inventory sheet's own Receipt/Consumption/Return/Current
                Stock figures disagree with its own log sheets (Current Stock
                is a static value in this workbook, so it can go stale)
  opening       the Opening Stock differs
  unexplained   whatever the causes above do not account for — reported, never
                hidden (a SKIP is not a PASS — rule 16)

⚠️ IT CHANGES NOTHING. The diagnosis reuses the sync's own planner
(`bulk_import.plan_ledger`) READ-ONLY, so "is this line in the workbook?" is
answered by exactly the matching the sync itself uses. Fixing is a person's
decision: add the line to the workbook, or have an admin void a test entry.

⚠️ THE ORIGINAL WORKBOOK IS NEVER WRITTEN. `mark_workbook()` returns a marked
COPY (rows filled, notes on the cells, a "GI Hub check" sheet): openpyxl drops
parts of this workbook it cannot read (its data-validation extensions), so
saving over the operator's file would damage it.
"""
from __future__ import annotations

import io
import json
from collections import defaultdict
from datetime import date, datetime
from typing import Optional

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

EPS = 1e-6
_KIND_SIGN = {"receipts": 1.0, "consumption": -1.0, "returns": -1.0}
_KIND_SHEET = {"receipts": "Receipt Log", "consumption": "Consumption Log",
               "returns": "Return Log"}
_KIND_WORD = {"receipts": "receipt", "consumption": "issue", "returns": "return"}
_ART = {"receipts": "A", "consumption": "An", "returns": "A"}
# who entered an app row, per table
_WHO = {"receipts": '"Received_by"', "consumption": '"Issued_By"', "returns": "NULL"}


def _num(v) -> Optional[float]:
    try:
        return None if v is None or v == "" else float(v)
    except (TypeError, ValueError):
        return None


def _s(v) -> str:
    return str(v).strip() if v is not None else ""


def _day(v) -> str:
    if isinstance(v, (datetime, date)):
        return v.isoformat()[:10]
    return _s(v)[:10]


def _fmt(q) -> str:
    q = float(q or 0)
    return f"{q:g}" if abs(q - round(q, 4)) < EPS else f"{q:.4f}"


# ── reading the workbook, with Excel row numbers ─────────────────────────────
def _sheet(wb, name: str, probe: tuple[str, ...]):
    """(header→col index, [(excel_row, row values)]) for one sheet."""
    ws = next((wb[n] for n in wb.sheetnames if n.strip().lower() == name.lower()), None)
    if ws is None:
        return None, []
    rows = list(ws.iter_rows(values_only=True))
    for i, row in enumerate(rows[:6]):
        cells = [_s(c).lower() for c in row]
        if all(p in cells for p in probe):
            head = {c: j for j, c in enumerate(cells) if c}
            return head, [(i + 2 + k, r) for k, r in enumerate(rows[i + 1:])]
    return None, []


def read_workbook(data: bytes) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        head, rows = _sheet(wb, "Inventory", ("sap code", "current stock"))
        inv: dict[str, dict] = {}
        if head:
            g = lambda r, k: r[head[k]] if k in head and head[k] < len(r) else None  # noqa: E731
            for xr, r in rows:
                sap = _s(g(r, "sap code"))
                if not sap:
                    continue
                inv[sap] = {"row": xr, "description": _s(g(r, "equipment description")),
                            "uom": _s(g(r, "uom")),
                            "opening": _num(g(r, "opening stock")) or 0.0,
                            "receipt": _num(g(r, "receipt")), "consumption": _num(g(r, "consumption")),
                            "return": _num(g(r, "return")), "current": _num(g(r, "current stock"))}
        logs: dict[str, list[dict]] = {}
        refcol = {"receipts": "dn. no.", "consumption": "tank no.", "returns": "reason"}
        for kind, name in _KIND_SHEET.items():
            h, rs = _sheet(wb, name, ("sap code", "qty."))
            out = []
            if h:
                for xr, r in rs:
                    sap = _s(r[h["sap code"]]) if h["sap code"] < len(r) else ""
                    qty = _num(r[h["qty."]]) if h["qty."] < len(r) else None
                    if not sap or qty is None:
                        continue
                    di = h.get("date")
                    ri = h.get(refcol[kind])
                    out.append({"row": xr, "sap": sap, "qty": qty,
                                "date": _day(r[di]) if di is not None and di < len(r) else "",
                                "ref": _s(r[ri]) if ri is not None and ri < len(r) else ""})
            logs[kind] = out
    finally:
        wb.close()
    return {"inventory": inv, "logs": logs}


# ── GI Hub's side ────────────────────────────────────────────────────────────
async def app_stock(session: AsyncSession) -> dict[str, float]:
    """Exactly the sync's verification formula (all sites, per SAP)."""
    return {r[0]: float(r[1]) for r in (await session.execute(text('''
        SELECT i."SAP_Code",
               COALESCE(i."Opening_Stock",0)
             + COALESCE((SELECT SUM(r."Quantity") FROM receipts r WHERE r."SAP_Code"=i."SAP_Code"),0)
             - COALESCE((SELECT SUM(c."Quantity") FROM consumption c WHERE c."SAP_Code"=i."SAP_Code"),0)
             - COALESCE((SELECT SUM(t."Quantity") FROM returns t WHERE t."SAP_Code"=i."SAP_Code"),0)
        FROM inventory i'''))).all()}


async def _opening(session: AsyncSession) -> dict[str, float]:
    return {r[0]: float(r[1] or 0) for r in (await session.execute(text(
        'SELECT "SAP_Code", "Opening_Stock" FROM inventory'))).all()}


async def _row_details(session: AsyncSession, kind: str, ids: list[int]) -> dict[int, dict]:
    if not ids:
        return {}
    rows = (await session.execute(text(
        f'SELECT id, "Remarks" AS remarks, {_WHO[kind]} AS who FROM {kind} WHERE id = ANY(:i)'),
        {"i": ids})).mappings().all()
    return {int(r["id"]): {"remarks": _s(r["remarks"]), "who": _s(r["who"])} for r in rows}


# ── the diagnosis ────────────────────────────────────────────────────────────
async def diagnose(session: AsyncSession, data: bytes, *, site_id: str) -> dict:
    """Read-only. Returns {total, matched, items:[{sap, workbook, app, difference,
    causes:[{code, effect, what, fix, where}], unexplained}]}."""
    from .. import bulk_import as bi
    book = read_workbook(data)
    inv, logs = book["inventory"], book["logs"]
    stock = await app_stock(session)
    opening = await _opening(session)
    mism = [s for s, v in inv.items() if v["current"] is not None
            and (s not in stock or abs(stock[s] - v["current"]) > EPS)]
    items: list[dict] = []
    if not mism:
        return {"total": sum(1 for v in inv.values() if v["current"] is not None),
                "matched": sum(1 for v in inv.values() if v["current"] is not None),
                "items": []}

    # the sync's own planner, read-only, inside a savepoint that is rolled back
    async with session.begin_nested() as sp:
        plan = await bi.plan_ledger(session, data, site_id)
        await sp.rollback()
    want = set(mism)
    by_sap: dict[str, list[dict]] = defaultdict(list)

    for kind, sec in plan["sections"].items():
        sign = _KIND_SIGN[kind]
        word = _KIND_WORD[kind]
        sheet = _KIND_SHEET[kind]
        mine = [r for r in sec.get("db_only_rows", []) if _s(r["sap"]) in want]
        det = await _row_details(session, kind, [int(r["id"]) for r in mine])
        for r in mine:
            d = det.get(int(r["id"]), {})
            src = _s(r.get("Source_Ref"))
            qr = src.startswith("SME_EXEC:")
            who = f" by {d['who']}" if d.get("who") else ""
            rem = f" (remarks: “{d['remarks']}”)" if d.get("remarks") else ""
            by_sap[_s(r["sap"])].append({
                "code": "app_only", "effect": sign * float(r["qty"] or 0),
                "what": (f"{_ART[kind]} {word} of {_fmt(r['qty'])} on {r['date']} was posted from a QR "
                         f"consumption form, and the {sheet} does not have it yet."
                         if qr else
                         f"{_ART[kind]} {word} of {_fmt(r['qty'])} on {r['date']} was entered in GI Hub"
                         f"{who}{rem}, and the {sheet} does not have it."),
                "fix": (f"If it really happened, add it to the {sheet}. If it was a test "
                        f"or a mistake, reverse it in GI Hub with a Stock Adjustment "
                        f"(Adjust Stock), or have an admin remove {kind} #{r['id']}."
                        if not qr else f"Add it to the {sheet} (see SME → QR ⇄ Excel)."),
                "where": {"ledger": kind, "id": int(r["id"])}})
        for r in sec.get("vanished", []):
            if _s(r["sap"]) not in want:
                continue
            by_sap[_s(r["sap"])].append({
                "code": "vanished", "effect": sign * float(r["qty"] or 0),
                "what": (f"GI Hub has {_ART[kind].lower()} {word} of {_fmt(r['qty'])} on {r['date']}"
                         f"{' for ' + r['ref'] if r.get('ref') else ''} that came from the "
                         f"{sheet}, but the {sheet} no longer has that line."),
                "fix": (f"If the line was deleted or re-dated in Excel by mistake, put it "
                        f"back. If the deletion is right, an admin removes it from GI Hub "
                        f"with the sync's --prune-vanished (it only deletes rows nothing "
                        f"else refers to)."),
                "where": {"ledger": kind, "id": int(r["id"])}})
        for r in sec.get("inserts", []):
            sap = _s(r.get("SAP_Code"))
            if sap not in want or str(r.get("Source_Ref", "")).endswith(":D"):
                continue
            row = next((x["row"] for x in logs.get(kind, []) if x["sap"] == sap
                        and x["date"] == _day(r.get("Date"))
                        and abs(x["qty"] - float(r["Quantity"])) < EPS), None)
            by_sap[sap].append({
                "code": "not_synced", "effect": -sign * float(r["Quantity"] or 0),
                "what": (f"The {sheet} has {_ART[kind].lower()} {word} of {_fmt(r['Quantity'])} on "
                         f"{_day(r.get('Date'))} that GI Hub does not have yet."),
                "fix": "Run the Excel sync — it will bring this line in.",
                "where": {"sheet": sheet, "row": row}})
        for r in sec.get("corrections", []):
            sap = _s(r.get("sap"))
            if sap not in want:
                continue
            delta = float(r["qty_from"] or 0) - float(r["qty_to"] or 0)
            by_sap[sap].append({
                "code": "qty_edited", "effect": sign * delta,
                "what": (f"{_ART[kind]} {word} on {r['date']} was changed in the {sheet} from "
                         f"{_fmt(r['qty_from'])} to {_fmt(r['qty_to'])} after GI Hub "
                         f"last synced."),
                "fix": "Run the Excel sync — it updates the line in place.",
                "where": {"ledger": kind, "id": int(r["id"])}})

    log_sum: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for kind, rows in logs.items():
        for x in rows:
            log_sum[x["sap"]][kind] += x["qty"]

    for sap in mism:
        v = inv[sap]
        app = stock.get(sap)
        causes = by_sap.get(sap, [])
        if app is None:
            items.append({"sap": sap, "description": v["description"], "uom": v["uom"],
                          "excel_row": v["row"], "workbook": v["current"], "app": None,
                          "difference": None, "unexplained": 0.0, "causes": [{
                              "code": "missing", "effect": 0.0,
                              "what": "This SAP is in the workbook but not in GI Hub's inventory.",
                              "fix": "Run the Excel sync (inventory) to add it.",
                              "where": {"sheet": "Inventory", "row": v["row"]}}]})
            continue
        ls = log_sum[sap]
        from_logs = v["opening"] + ls["receipts"] - ls["consumption"] - ls["returns"]
        if abs(from_logs - v["current"]) > EPS:
            causes.append({
                "code": "sheet_totals", "effect": from_logs - v["current"],
                "what": (f"The Inventory sheet says Current Stock {_fmt(v['current'])}, but "
                         f"its own logs add up to {_fmt(from_logs)} (Opening "
                         f"{_fmt(v['opening'])} + Receipts {_fmt(ls['receipts'])} − Issues "
                         f"{_fmt(ls['consumption'])} − Returns {_fmt(ls['returns'])})."),
                "fix": "Correct the Inventory sheet's figures (Current Stock is a static "
                       "value in this workbook, so it goes stale when a log changes).",
                "where": {"sheet": "Inventory", "row": v["row"]}})
        d_open = opening.get(sap, 0.0) - v["opening"]
        if abs(d_open) > EPS:
            causes.append({
                "code": "opening", "effect": d_open,
                "what": (f"Opening Stock is {_fmt(opening.get(sap, 0.0))} in GI Hub and "
                         f"{_fmt(v['opening'])} in the workbook."),
                "fix": "Run the Excel sync (inventory), or correct the Opening Stock in Excel.",
                "where": {"sheet": "Inventory", "row": v["row"]}})
        diff = app - v["current"]
        unexplained = round(diff - sum(c["effect"] for c in causes), 6)
        if abs(unexplained) > EPS:
            causes.append({
                "code": "unexplained", "effect": unexplained,
                "what": f"{_fmt(unexplained)} of the difference is not explained by any line above.",
                "fix": "Compare this material's rows in GI Hub (Records) with the logs by hand.",
                "where": {"sheet": "Inventory", "row": v["row"]}})
        for c in causes:
            c["effect"] = round(c["effect"], 6)
        items.append({"sap": sap, "description": v["description"], "uom": v["uom"],
                      "excel_row": v["row"], "workbook": v["current"], "app": round(app, 6),
                      "difference": round(diff, 6),
                      "unexplained": unexplained if abs(unexplained) > EPS else 0.0,
                      "causes": causes})
    total = sum(1 for v in inv.values() if v["current"] is not None)
    items.sort(key=lambda x: (x["sap"]))
    return {"total": total, "matched": total - len(items), "items": items}


# ── storing and reading the latest check ─────────────────────────────────────
async def store(session: AsyncSession, result: dict, *, site_id: str,
                workbook: str, username: str) -> int:
    return (await session.execute(text(
        'INSERT INTO stock_excel_checks ("Site_ID", workbook, checked_by, total_saps, '
        'matched, items) VALUES (:s, :w, :u, :t, :m, :i) RETURNING id'),
        {"s": site_id, "w": workbook, "u": username, "t": result["total"],
         "m": result["matched"], "i": json.dumps(result["items"])})).scalar_one()


async def latest(session: AsyncSession, *, site_id: Optional[str]) -> Optional[dict]:
    q = ('SELECT id, "Site_ID", workbook, checked_by, checked_at, total_saps, matched, items '
         'FROM stock_excel_checks')
    params: dict = {}
    if site_id is not None:
        q += ' WHERE "Site_ID" = :s'
        params["s"] = site_id
    r = (await session.execute(text(q + " ORDER BY checked_at DESC, id DESC LIMIT 1"),
                               params)).mappings().first()
    if r is None:
        return None
    d = dict(r)
    d["items"] = json.loads(d["items"] or "[]")
    d["checked_at"] = d["checked_at"].isoformat() if d["checked_at"] else None
    return d


# ── the marked COPY of the workbook ──────────────────────────────────────────
RED = "FFF4CCCC"
AMBER = "FFFCE8B2"
_HEAD = ["SAP", "Description", "Workbook stock", "GI Hub stock", "Difference",
         "What happened", "How to fix", "Where"]


def mark_workbook(data: bytes, result: dict, *, checked_at: str = "") -> bytes:
    """A COPY of the workbook with every mismatch marked. Never the original."""
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.load_workbook(io.BytesIO(data))
    red = PatternFill("solid", fgColor=RED)
    amber = PatternFill("solid", fgColor=AMBER)
    inv = next((wb[n] for n in wb.sheetnames if n.strip().lower() == "inventory"), None)
    head_row, cur_col = None, None
    if inv is not None:
        for r in range(1, 7):
            for c in range(1, inv.max_column + 1):
                v = _s(inv.cell(r, c).value).lower()
                if v == "current stock":
                    head_row, cur_col = r, c
    for it in result["items"]:
        if inv is None or not it.get("excel_row"):
            continue
        row = it["excel_row"]
        for c in range(1, inv.max_column + 1):
            inv.cell(row, c).fill = red
        if cur_col:
            lines = [f"GI Hub: {_fmt(it['app']) if it['app'] is not None else '—'} · "
                     f"this sheet: {_fmt(it['workbook'])}"]
            lines += [f"• {c['what']}" for c in it["causes"]][:6]
            lines.append("Details: sheet “GI Hub check”.")
            cm = Comment("\n".join(lines), "GI Hub")
            cm.width, cm.height = 420, 60 + 36 * len(lines)
            inv.cell(row, cur_col).comment = cm
        # the log lines the causes point at
        for c in it["causes"]:
            w = c.get("where") or {}
            if w.get("sheet") and w.get("row") and w["sheet"] != "Inventory":
                ws = next((wb[n] for n in wb.sheetnames
                           if n.strip().lower() == w["sheet"].lower()), None)
                if ws is not None:
                    for col in range(1, min(ws.max_column, 20) + 1):
                        ws.cell(w["row"], col).fill = amber
                    ws.cell(w["row"], 1).comment = Comment(f"{c['what']}\n{c['fix']}", "GI Hub")

    name = "GI Hub check"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name, 0)
    ws["A1"] = (f"GI Hub vs this workbook — {result['matched']} of {result['total']} "
                f"materials match" + (f" (checked {checked_at[:16].replace('T', ' ')})"
                                      if checked_at else ""))
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("This is a MARKED COPY. Make the fixes in your own workbook, then run the "
                "Excel sync again. Red rows on the Inventory sheet are the materials below; "
                "amber rows on a log sheet are lines named here.")
    ws["A2"].font = Font(italic=True, color="FF555555")
    for j, h in enumerate(_HEAD, 1):
        cell = ws.cell(4, j, h)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF1F3A5F")
    r = 5
    for it in result["items"]:
        first = True
        for c in it["causes"] or [{"what": "—", "fix": "—", "where": {}}]:
            w = c.get("where") or {}
            where = (f"{w['sheet']} row {w['row']}" if w.get("sheet") and w.get("row")
                     else f"GI Hub {w['ledger']} #{w['id']}" if w.get("ledger") else "")
            vals = [it["sap"] if first else "", it["description"] if first else "",
                    it["workbook"] if first else None, it["app"] if first else None,
                    it["difference"] if first else None, c["what"], c["fix"], where]
            for j, v in enumerate(vals, 1):
                cell = ws.cell(r, j, v)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if first:
                    cell.fill = red
            if first and inv is not None and it.get("excel_row"):
                ws.cell(r, 1).hyperlink = f"#'{inv.title}'!A{it['excel_row']}"
                ws.cell(r, 1).font = Font(color="FF0B5394", underline="single")
            first = False
            r += 1
    for col, width in zip("ABCDEFGH", (10, 32, 14, 14, 12, 70, 60, 22)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
