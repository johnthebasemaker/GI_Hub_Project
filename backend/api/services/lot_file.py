"""
backend/api/services/lot_file.py — Phase 16b: the Lot Register workbook.

`Rubber & Brick Materials - CNCEC.xlsx` keeps, per Surface Shield material,
one sheet of what arrived: batch, MFD, expiry, DN, quantity. GI Hub reads it to
DESCRIBE its lots — and never to count them.

⚠️ THE SAFETY PRINCIPLE (plan §0, approved 2026-10-01): THIS FILE NEVER MOVES
STOCK. It writes no ledger quantity. It writes, and only writes:

    lots        MFD, expiry (+ where the expiry came from), batch reference,
                DN — on the lot the ledger already has, or a new lot row for a
                batch the receipts did not name
    lot_units   the CHEMOLINE roll register: roll → batch (ruling Q16-3)
    receipts    `Lot_Number` on an EXISTING receipt that has none, when the
                file names exactly one batch for that receipt
    inventory   `Shelf_Life_Months`, only where unset, learned from the
                file's own MFD → expiry pairs (ruling Q16-5)

Its quantities are CHECKED against the Receipt Log (same SAP + DN + received
date) and every disagreement is reported. A mistake in this file can make an
expiry wrong; it can never double a quantity.

THE THREE LAYOUTS (plan §2.1), told apart by their headers:

    batch    `Batch No.`  — PU A–D, Furan, Phenacin, Coroflake, BC 3004 …
    pallet   `Package No.`, no batch — the bricks (not lot-tracked, Q16-4)
    roll     `Roll No.` + `Order No.` — CHEMOLINE: the roll is a unit, the
             Order No. is its batch

THE KEY. Each row's SAP comes from the `SAP` column (ruling Q16-14 — the
operator added it to every sheet). Without it, a deterministic resolver
matches the material name, the component letter (COMP C) and the pack size
(10 Kg vs 2.5 Kg tells 1041-3 from 1041-4) and refuses anything ambiguous —
`Material_Code` is never the key (COMP A–D share one), only a cross-check.
"""
from __future__ import annotations

import datetime as _dt
import re
import statistics
from collections import defaultdict
from typing import Any, Optional

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from . import lots as LOTS

# Header spellings seen in the real workbook (2026-10-01), lower-cased.
_H = {
    "sap": ("sap", "sap code", "sap_code"),
    "material": ("material",),
    "batch": ("batch no.", "batch no", "batch no. & mfg. date", "lot", "lot no."),
    "roll": ("roll no.", "roll no"),
    "order": ("order no.", "order no"),
    "package": ("package no.", "package no"),
    "size": ("size of package", "size"),
    "mfd": ("mfg. date", "mfd", "mfg date"),
    "qty": ("quantity", "qty", "received qty"),
    "received": ("received date", "rcvd date"),
    "expiry": ("expiry date", "exp. date", "expiry"),
    "dn": ("dn no.", "dn no", "dn. no."),
    "code": ("materil code", "material code", "material name"),
    "pallet": ("pallet",),
    "location": ("location",),
    "sqm": ("sqm",),
}


def _hkey(v: Any) -> str:
    return " ".join(str(v or "").split()).lower()


def _date(v: Any) -> Optional[str]:
    """A real date → 'YYYY-MM-DD'. Text such as '(Aug)-2026' is not a date."""
    if isinstance(v, _dt.datetime):
        return v.date().isoformat()
    if isinstance(v, _dt.date):
        return v.isoformat()
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return _dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _num(v: Any) -> Optional[float]:
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"(\d+(?:\.\d+)?)", str(v or ""))
    return float(m.group(1)) if m else None


def _text(v: Any) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = " ".join(str(v).split())
    return s or None


def add_months(day: str, months: int) -> str:
    d = _dt.date.fromisoformat(day)
    y, m = divmod(d.month - 1 + months, 12)
    y, m = d.year + y, m + 1
    last = [31, 29 if (y % 4 == 0 and (y % 100 or y % 400 == 0)) else 28, 31, 30, 31,
            30, 31, 31, 30, 31, 30, 31][m - 1]
    return _dt.date(y, m, min(d.day, last)).isoformat()


def months_between(a: str, b: str) -> int:
    da, db = _dt.date.fromisoformat(a), _dt.date.fromisoformat(b)
    return round((db - da).days / 30.4375)


# ── reading ──────────────────────────────────────────────────────────────────
def read_workbook(data: bytes) -> dict:
    """Every sheet → typed rows. Pure; no database. Header = the first row
    within the top five that carries a MATERIAL column."""
    import io

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheets, rows_out, notes = [], [], []
    for ws in wb.worksheets:
        top = []
        for r in ws.iter_rows(values_only=True):
            top.append(r)
            if len(top) >= 5:
                break
        hi = next((i for i, r in enumerate(top) if r and any(_hkey(c) == "material" for c in r)), None)
        if hi is None:
            notes.append(f"{ws.title!r}: no MATERIAL header in the first five rows — skipped")
            continue
        hdr = [_hkey(c) for c in top[hi]]
        col = {k: next((i for i, h in enumerate(hdr) if h in names), None)
               for k, names in _H.items()}
        layout = ("roll" if col["roll"] is not None else
                  "pallet" if col["batch"] is None and col["package"] is not None else "batch")
        sheets.append({"sheet": ws.title.strip(), "layout": layout,
                       "has_sap": col["sap"] is not None})

        def g(r, k):
            i = col[k]
            return r[i] if i is not None and i < len(r) else None

        for n, r in enumerate(ws.iter_rows(min_row=hi + 2, values_only=True), start=hi + 2):
            if not r or not any(c not in (None, "") for c in r):
                continue
            material = _text(g(r, "material"))
            sap = _text(g(r, "sap"))
            row = {"sheet": ws.title.strip(), "row": n, "layout": layout,
                   "sap": sap.upper() if sap else None, "material": material,
                   "size": _text(g(r, "size")), "size_n": _num(g(r, "size")),
                   "qty": _num(g(r, "qty")), "received": _date(g(r, "received")),
                   "mfd": _date(g(r, "mfd")), "mfd_raw": _text(g(r, "mfd")),
                   "expiry": _date(g(r, "expiry")), "dn": _text(g(r, "dn")),
                   "code": _text(g(r, "code")), "pallet": _text(g(r, "pallet")),
                   "location": _text(g(r, "location")), "sqm": _text(g(r, "sqm"))}
            if layout == "roll":
                row["unit"] = LOTS.norm_roll(g(r, "roll"))
                row["lot"] = LOTS.norm_lot(g(r, "order")) or LOTS.roll_batch(row["unit"])
                if row["lot"] and row["lot"].startswith("10"):   # same O/0 slip (Q16-12)
                    row["lot"] = "1O" + row["lot"][2:]
            elif layout == "batch":
                row["lot"] = LOTS.norm_lot(g(r, "batch"))
            else:
                row["lot"] = None
            rows_out.append(row)
    return {"sheets": sheets, "rows": rows_out, "notes": notes}


# ── the SAP, when a row has none ─────────────────────────────────────────────
def _norm_name(s: Optional[str]) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(s or "").upper())


def resolve_sap(row: dict, inventory: dict[str, dict]) -> tuple[Optional[str], str]:
    """Name + component letter + pack size → exactly one SAP, or a refusal."""
    name = str(row.get("material") or "").upper()
    if not name:
        return None, "no material name"
    words = [w for w in re.findall(r"[A-Z0-9]{3,}", name) if w not in ("COMP", "KG")]
    comp = re.search(r"COMP\s*([A-D])\b", name)
    thick = re.search(r"\((\d+)\s*MM\)", name)
    cands = []
    for sap, it in inventory.items():
        desc = str(it.get("desc") or "").upper()
        if comp:
            dcomp = re.search(r"COMP\s*([A-D])\b", desc)
            if not dcomp or dcomp.group(1) != comp.group(1):
                continue
        if thick and f"({thick.group(1)}MM)" not in desc.replace(" ", ""):
            continue
        score = sum(1 for w in words if w in desc)
        if score >= max(1, len(words) - 1):
            cands.append((score, sap, it))
    if not cands:
        return None, "no inventory item matches the material name"
    best = max(c[0] for c in cands)
    top = [c for c in cands if c[0] == best]
    if len(top) > 1 and row.get("size_n") is not None:
        top = [c for c in top if c[2].get("unit_size") is not None
               and abs(float(c[2]["unit_size"]) - row["size_n"]) < 0.05] or top
    if len(top) == 1:
        return top[0][1], "name"
    return None, f"ambiguous: {', '.join(sorted(c[1] for c in top))} — add the SAP column"


# ── planning ─────────────────────────────────────────────────────────────────
async def plan(session: AsyncSession, data: bytes, site_id: str) -> dict:
    """Read the file and decide what to write. Writes nothing."""
    book = read_workbook(data)
    inv = {r[0]: {"desc": r[1], "code": r[2], "unit_size": r[3], "life": r[4]}
           for r in (await session.execute(text(
               'SELECT TRIM("SAP_Code"), "Equipment_Description", "Material_Code", '
               '"Unit_Size", "Shelf_Life_Months" FROM inventory'))).all()}
    track = await LOTS.tracking(session)
    out: dict = {"site_id": site_id, "sheets": book["sheets"], "rows": len(book["rows"]),
                 "resolved": 0, "matched": 0, "lot_rows": 0, "rejects": [],
                 "warnings": list(book["notes"]), "qty_mismatch": [], "unmatched": [],
                 "no_batch": [], "expiry_refused": [], "date_conflicts": [],
                 "code_mismatch": [], "orphans": [], "lots": [], "lot_changes": [],
                 "file_only": [],
                 "units": [], "receipt_fills": [], "shelf_life": {}}

    # receipts of this site, by (SAP, day) and (SAP, DN, day)
    recs = [dict(r) for r in (await session.execute(text(
        'SELECT id, TRIM("SAP_Code") AS sap, SUBSTRING("Date" FROM 1 FOR 10) AS day, '
        '"Quantity" AS qty, TRIM(COALESCE("DN_No", "DN_Number", \'\')) AS dn, '
        '"Lot_Number" AS lot FROM receipts WHERE "Site_ID" = :s'), {"s": site_id})).mappings().all()]
    by_dn = defaultdict(list)
    by_day = defaultdict(list)
    for r in recs:
        by_dn[(r["sap"], r["dn"], r["day"])].append(r)
        by_day[(r["sap"], r["day"])].append(r)

    file_qty: dict[tuple, float] = defaultdict(float)
    lot_rows: dict[tuple, list[dict]] = defaultdict(list)
    fill_groups: dict[tuple, set] = defaultdict(set)
    life_obs: dict[str, list[int]] = defaultdict(list)
    for row in book["rows"]:
        sap = row["sap"]
        if not row["material"] and not sap:
            if row.get("lot") or row.get("mfd") or row.get("expiry"):
                out["orphans"].append({"sheet": row["sheet"], "row": row["row"],
                                       "lot": row.get("lot")})
            continue
        if sap is None:
            sap, how = resolve_sap(row, inv)
            if sap is None:
                out["rejects"].append({"sheet": row["sheet"], "row": row["row"], "reason": how})
                continue
        if sap not in inv:
            out["rejects"].append({"sheet": row["sheet"], "row": row["row"],
                                   "reason": f"SAP {sap} is not in the inventory master"})
            continue
        out["resolved"] += 1
        row["sap"] = sap
        # the code is a check, never the key — a variant SAP carries none, so
        # compare with its family's (1042-3 → 1042)
        fam = inv.get(sap.split("-")[0], {}).get("code") or inv[sap].get("code")
        if row["code"] and fam and row["code"].upper() != str(fam).upper():
            out["code_mismatch"].append({"sheet": row["sheet"], "row": row["row"], "sap": sap,
                                         "file": row["code"], "inventory": fam})
        # cross-check against the Receipt Log
        if row["layout"] == "roll":
            key, match = (sap, row["received"]), by_day.get((sap, row["received"]), [])
        else:
            key, match = (sap, row["dn"] or "", row["received"]), \
                by_dn.get((sap, row["dn"] or "", row["received"]), [])
        file_qty[key] += 1.0 if row["layout"] == "roll" else float(row["qty"] or 0)
        if match:
            out["matched"] += 1
        else:
            out["unmatched"].append({"sheet": row["sheet"], "row": row["row"], "sap": sap,
                                     "dn": row["dn"], "received": row["received"]})
        if track.get(sap) is None:          # the bricks: cross-check only (Q16-4)
            continue
        if not row.get("lot"):
            out["no_batch"].append({"sheet": row["sheet"], "row": row["row"], "sap": sap,
                                    "dn": row["dn"], "qty": row["qty"]})
            continue
        if row["expiry"] and row["mfd"] and row["expiry"] < row["mfd"]:
            out["expiry_refused"].append({"sheet": row["sheet"], "row": row["row"],
                                          "sap": sap, "lot": row["lot"], "mfd": row["mfd"],
                                          "expiry": row["expiry"]})
            row["expiry"] = None
        if row["mfd"] and row["expiry"]:
            life_obs[sap].append(months_between(row["mfd"], row["expiry"]))
        out["lot_rows"] += 1
        lot_rows[(row["lot"], sap)].append(row)
        if match and row["layout"] == "batch":
            for r in match:
                fill_groups[r["id"]].add(row["lot"])
        if row["layout"] == "roll" and row.get("unit"):
            out["units"].append({"Unit_No": row["unit"], "Lot_Number": row["lot"],
                                 "SAP_Code": sap, "Site_ID": site_id,
                                 "Received_Date": row["received"], "SQM": row["sqm"],
                                 "Pallet": row["pallet"], "Location": row["location"]})

    # quantities: the file against the Receipt Log, per receipt group
    for key, fq in file_qty.items():
        rq = sum(float(r["qty"] or 0) for r in
                 (by_day.get(key) if len(key) == 2 else by_dn.get(key)) or [])
        if (by_day.get(key) if len(key) == 2 else by_dn.get(key)) and abs(fq - rq) > 1e-6:
            out["qty_mismatch"].append({"key": key, "file": fq, "receipts": rq})

    # shelf life, learned only where the item has none (Q16-5)
    for sap, obs in life_obs.items():
        if inv[sap].get("life") is None and obs:
            out["shelf_life"][sap] = int(statistics.median(obs))
    life = {sap: (inv[sap].get("life") or out["shelf_life"].get(sap)) for sap in inv}

    # one row per lot — the earliest expiry wins when rows disagree (Q16-6)
    existing = {(r[0], r[1]): dict(zip(("lot", "sap", "mfd", "exp", "src", "batch", "dn",
                                        "received", "source"), r))
                for r in (await session.execute(text(
                    'SELECT "Lot_Number", "SAP_Code", "MFD_Date", "Expiry_Date", '
                    '"Expiry_Source", "Batch_Ref", "DN_No", "Received_Date", "Source" '
                    'FROM lots WHERE "Site_ID" = :s'), {"s": site_id})).all()}
    for (lot, sap), rs in sorted(lot_rows.items()):
        exps = sorted({r["expiry"] for r in rs if r["expiry"]})
        mfds = sorted({r["mfd"] for r in rs if r["mfd"]})
        # rolls of one batch are made on different days — only an EXPIRY
        # disagreement means something there
        if len(exps) > 1 or (len(mfds) > 1 and rs[0]["layout"] != "roll"):
            out["date_conflicts"].append({"sap": sap, "lot": lot, "mfd": mfds, "expiry": exps})
        mfd = mfds[0] if mfds else None
        if exps:
            expiry, src = exps[0], "file"
        elif mfd and life.get(sap):
            expiry, src = add_months(mfd, int(life[sap])), "derived"
        else:
            expiry, src = None, None
        received = min((r["received"] for r in rs if r["received"]), default=None)
        want = {"MFD_Date": mfd, "Expiry_Date": expiry, "Expiry_Source": src,
                "Batch_Ref": rs[0]["lot"] if rs[0]["layout"] == "roll" else None,
                "DN_No": next((r["dn"] for r in rs if r["dn"]), None),
                "Received_Date": received}
        cur = existing.get((lot, sap))
        if cur and cur.get("src") == "app":
            # a person typed this expiry on the Receive form — the app wins
            want.pop("Expiry_Date")
            want.pop("Expiry_Source")
        if cur is None:
            out["lots"].append({"Lot_Number": lot, "SAP_Code": sap, "Site_ID": site_id, **want})
            # A batch lot no receipt names: usually the Receipt Log spells it
            # differently ('0.926' vs '0926'). Rolls are expected here — their
            # receipts carry no roll numbers, the register is their record.
            if rs[0]["layout"] == "batch":
                named = sorted({r["lot"] for r in recs if r["sap"] == sap and r["lot"]})
                out["file_only"].append({"sap": sap, "lot": lot, "receipt_lots": named})
        else:
            cmp = {"MFD_Date": cur["mfd"], "Expiry_Date": cur["exp"],
                   "Expiry_Source": cur["src"], "Batch_Ref": cur["batch"],
                   "DN_No": cur["dn"], "Received_Date": cur["received"]}
            diff = {k: v for k, v in want.items()
                    if v is not None and str(cmp.get(k) or "") != str(v)
                    and not (k == "Received_Date" and cmp.get(k) and cmp[k] <= v)}
            if diff:
                out["lot_changes"].append({"Lot_Number": lot, "SAP_Code": sap,
                                           "Site_ID": site_id, "diff": diff})

    # a receipt with no lot, whose file rows name exactly one batch
    rec_by_id = {r["id"]: r for r in recs}
    for rid, lots_named in fill_groups.items():
        r = rec_by_id[rid]
        if not r["lot"] and len(lots_named) == 1:
            out["receipt_fills"].append({"id": rid, "sap": r["sap"], "day": r["day"],
                                         "dn": r["dn"], "Lot_Number": next(iter(lots_named))})

    # roll register: only what is new or changed
    have = {(r[0], r[1]): r[2] for r in (await session.execute(text(
        'SELECT "SAP_Code", "Unit_No", "Lot_Number" FROM lot_units')))}
    out["units_new"] = sum(1 for u in out["units"] if (u["SAP_Code"], u["Unit_No"]) not in have)
    out["units_changed"] = sum(1 for u in out["units"]
                               if have.get((u["SAP_Code"], u["Unit_No"])) not in (None, u["Lot_Number"]))
    return out


# ── applying ─────────────────────────────────────────────────────────────────
async def apply(session: AsyncSession, p: dict, username: str) -> dict:
    """Write the plan's DESCRIPTIONS. No ledger quantity is touched."""
    from .ledger import write_audit
    site = p["site_id"]
    for lot in p["lots"]:
        await session.execute(text('''
            INSERT INTO lots ("Lot_Number", "SAP_Code", "Site_ID", "Received_Date",
                              "Expiry_Date", "Expiry_Source", "MFD_Date", "Batch_Ref",
                              "DN_No", "Status", "Source", "updated_at")
            VALUES (:Lot_Number, :SAP_Code, :Site_ID, COALESCE(:Received_Date, ''),
                    :Expiry_Date, :Expiry_Source, :MFD_Date, :Batch_Ref, :DN_No,
                    'open', 'lotfile', now())
            ON CONFLICT ("Lot_Number", "SAP_Code", "Site_ID") DO NOTHING'''),
            {k: lot.get(k) for k in ("Lot_Number", "SAP_Code", "Site_ID", "Received_Date",
                                     "Expiry_Date", "Expiry_Source", "MFD_Date",
                                     "Batch_Ref", "DN_No")})
    for ch in p["lot_changes"]:
        sets = ", ".join(f'"{k}" = :{k}' for k in ch["diff"])
        await session.execute(text(
            f'UPDATE lots SET {sets}, "updated_at" = now() WHERE "Lot_Number" = :lot '
            f'AND "SAP_Code" = :sap AND "Site_ID" = :site'),
            {**ch["diff"], "lot": ch["Lot_Number"], "sap": ch["SAP_Code"], "site": site})
    for u in p["units"]:
        await session.execute(text('''
            INSERT INTO lot_units ("Unit_No", "Lot_Number", "SAP_Code", "Site_ID",
                                   "Received_Date", "SQM", "Pallet", "Location", "Source")
            VALUES (:Unit_No, :Lot_Number, :SAP_Code, :Site_ID, :Received_Date, :SQM,
                    :Pallet, :Location, 'lotfile')
            ON CONFLICT ("SAP_Code", "Unit_No") DO UPDATE SET
                "Lot_Number" = EXCLUDED."Lot_Number",
                "Received_Date" = COALESCE(EXCLUDED."Received_Date", lot_units."Received_Date"),
                "SQM" = COALESCE(EXCLUDED."SQM", lot_units."SQM"),
                "Pallet" = COALESCE(EXCLUDED."Pallet", lot_units."Pallet"),
                "Location" = COALESCE(EXCLUDED."Location", lot_units."Location"),
                "updated_at" = now()
            WHERE (lot_units."Lot_Number", lot_units."Received_Date", lot_units."SQM",
                   lot_units."Pallet", lot_units."Location")
                  IS DISTINCT FROM (EXCLUDED."Lot_Number",
                   COALESCE(EXCLUDED."Received_Date", lot_units."Received_Date"),
                   COALESCE(EXCLUDED."SQM", lot_units."SQM"),
                   COALESCE(EXCLUDED."Pallet", lot_units."Pallet"),
                   COALESCE(EXCLUDED."Location", lot_units."Location"))'''), u)
    for f in p["receipt_fills"]:
        await session.execute(text(
            'UPDATE receipts SET "Lot_Number" = :lot WHERE id = :id '
            'AND COALESCE("Lot_Number", \'\') = \'\''), {"lot": f["Lot_Number"], "id": f["id"]})
    for sap, months in p["shelf_life"].items():
        await session.execute(text(
            'UPDATE inventory SET "Shelf_Life_Months" = :m WHERE TRIM("SAP_Code") = :s '
            'AND "Shelf_Life_Months" IS NULL'), {"m": months, "s": sap})
    # receipts just given a lot become lots too (the 16a path, once more)
    made = await LOTS.sync_lots_from_ledger(session, site)
    if p["lots"] or p["lot_changes"] or p["units_new"] or p["receipt_fills"] or p["shelf_life"]:
        await write_audit(session, username, "LOT_FILE_SYNC", "lots",
                          f"+{len(p['lots'])} lot(s) ~{len(p['lot_changes'])} · "
                          f"+{p['units_new']} roll(s) · {len(p['receipt_fills'])} receipt "
                          f"lot(s) filled · shelf life set for {len(p['shelf_life'])}")
    return {"lots_new": len(p["lots"]), "lots_changed": len(p["lot_changes"]),
            "units_new": p["units_new"], "receipt_fills": len(p["receipt_fills"]),
            "shelf_life_set": len(p["shelf_life"]), "lots_from_receipts": made["created"]}
